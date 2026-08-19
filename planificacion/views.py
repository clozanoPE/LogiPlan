from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse
# Se sincroniza con models incluyendo CapacidadDiaria
from .models import PlanificacionDespacho, ConfiguracionCapacidad, CapacidadDiaria
from .forms import PlanificacionForm
from datetime import date, datetime, timedelta
import collections
from io import BytesIO
from django.utils import timezone
from django.db import models # Añadido para Q objects en búsqueda
from django.db.models import Count, Sum, Avg  # Agregaciones ORM para el módulo de Estadística
from django.core.paginator import Paginator
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --- CONSTANTES COMPARTIDAS ---
DIAS_SEMANA_ES = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
# Límite defensivo de rango para el módulo de Estadística: evita recorrer día a día
# rangos arbitrariamente grandes elegidos libremente por el usuario vía filtros GET.
MAX_RANGO_ESTADISTICAS_DIAS = 180
# Cantidad fija de registros por página en la vista "Consultar" (lista_busqueda).
REGISTROS_POR_PAGINA = 10

# --- CARGA MASIVA POR EXCEL ---
# Columnas esperadas en el archivo .xlsx de carga masiva (nombre exacto, orden no relevante).
COLUMNAS_EXCEL_CARGA_MASIVA = [
    'fecha', 'placa', 'conductor', 'cliente', 'sede_destino',
    'origen', 'orden_compra', 'pallets', 'costo_transporte',
]
# Límite defensivo de filas por archivo, para evitar cargas accidentalmente enormes.
MAX_FILAS_CARGA_MASIVA = 500
# Color corporativo Daryza, reutilizado en las cabeceras de todos los archivos Excel generados.
COLOR_HEADER_EXCEL = '2BB673'

# --- FUNCIONES DE AYUDA PARA GRUPOS ---

def es_planeamiento(user):
    """Verifica si el usuario pertenece al grupo Planeamiento o es Superusuario."""
    return user.groups.filter(name='Planeamiento').exists() or user.is_superuser

def obtener_nombre_dia_es(f):
    return DIAS_SEMANA_ES[f.weekday()]

def obtener_nombre_mes_es(f):
    meses = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
             'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
    return meses[f.month - 1]


def _parsear_fecha_aaaammdd(valor):
    """
    Parsea una celda de fecha en formato estricto AAAAMMDD (texto o número)
    a un objeto date. También admite que openpyxl la haya leído como
    date/datetime nativo (si Excel autoformateó la celda como fecha).
    Lanza ValueError con un mensaje legible si el formato no es válido.
    """
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor

    texto = str(valor).strip()
    if texto.endswith('.0'):  # celdas numéricas tipo 20260820.0
        texto = texto[:-2]
    if not (texto.isdigit() and len(texto) == 8):
        raise ValueError(f"'{valor}' no tiene el formato AAAAMMDD esperado")
    return datetime.strptime(texto, '%Y%m%d').date()


def _generar_excel_despachos(despachos, nombre_archivo):
    """
    Genera un archivo .xlsx en memoria (BytesIO) a partir de un queryset/lista
    de PlanificacionDespacho, con cabecera estilizada (verde Daryza) y tipos
    de dato correctos (fecha, entero, decimal). Retorna un HttpResponse listo
    para descargar. Compartido entre todas las exportaciones del sistema.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Despachos'

    encabezados = [
        'Fecha Despacho', 'Cliente', 'Sede', 'Placa', 'Conductor',
        'Orden de Compra', 'Origen', 'Pallets', 'Costo Transporte',
    ]
    relleno_header = PatternFill(start_color=COLOR_HEADER_EXCEL, end_color=COLOR_HEADER_EXCEL, fill_type='solid')
    fuente_header = Font(color='FFFFFF', bold=True)
    for col_num, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col_num, value=encabezado)
        celda.fill = relleno_header
        celda.font = fuente_header
        celda.alignment = Alignment(horizontal='center')

    for row_num, d in enumerate(despachos, start=2):
        ws.cell(row=row_num, column=1, value=d.fecha_despacho).number_format = 'DD/MM/YYYY'
        ws.cell(row=row_num, column=2, value=d.cliente)
        ws.cell(row=row_num, column=3, value=d.sede)
        ws.cell(row=row_num, column=4, value=d.unidad_placa)
        ws.cell(row=row_num, column=5, value=d.conductor)
        ws.cell(row=row_num, column=6, value=d.orden_compra)
        ws.cell(row=row_num, column=7, value=d.carga_origen)
        ws.cell(row=row_num, column=8, value=d.pallets).number_format = '#,##0'
        ws.cell(row=row_num, column=9, value=float(d.costo_transporte or 0)).number_format = '#,##0.00'

    for col_num, encabezado in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(col_num)].width = max(len(encabezado) + 4, 14)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def calcular_kpis_dashboard(lunes_actual, conteo_por_fecha, capacidad_limite_nominal, excepciones_externas=None):
    kpis = {
        'tendencia': [],
        'eficiencia': 0,
        'dia_pico': "N/A",
        'alertas': 0,
        'total_unidades': 0
    }
    
    # Prioridad: Excepciones externas (pasadas desde la vista dashboard)
    excepciones_semana = excepciones_externas if excepciones_externas is not None else {}

    max_unidades = -1
    total_unidades_semana = 0
    dias_criticos = 0
    capacidad_total_semana = 0
    nombres_cortos = ["LUN", "MAR", "MIE", "JUE", "VIE", "SAB"]

    for i in range(6): 
        fecha = lunes_actual + timedelta(days=i)
        # Obtenemos placas únicas para ese día
        unidades = len(conteo_por_fecha.get(fecha, set()))
        
        # LÓGICA CLAVE: Si hay excepción se usa, si no, el nominal (global)
        limite_real = excepciones_semana.get(fecha, capacidad_limite_nominal)
        capacidad_total_semana += limite_real
        
        porcentaje = (unidades / limite_real * 100) if limite_real > 0 else 0
        
        kpis['tendencia'].append({
            'dia': nombres_cortos[i],
            'valor': min(porcentaje, 100),
            'real': porcentaje,
            'unidades': unidades
        })

        total_unidades_semana += unidades
        if unidades > max_unidades:
            max_unidades = unidades
            kpis['dia_pico'] = obtener_nombre_dia_es(fecha)
        
        # ALERTA: Si se iguala o supera el límite real de ese día específico
        if unidades >= limite_real and limite_real > 0:
            dias_criticos += 1

    kpis['total_unidades'] = total_unidades_semana
    kpis['eficiencia'] = round((total_unidades_semana / capacidad_total_semana * 100), 1) if capacidad_total_semana > 0 else 0
    kpis['alertas'] = dias_criticos
    
    return kpis

# --- VISTAS DEL SISTEMA ---
@login_required
def dashboard(request):
    """
    Dashboard principal. Soporta navegación histórica de solo lectura por
    semanas vía ?week_offset=N (0 = semana actual, negativo = semanas pasadas).
    La ventana visual siempre cubre 14 días (Semana N + Semana N+1),
    desplazada según el offset.
    """
    config = ConfiguracionCapacidad.objects.first()
    capacidad_limite_global = config.limite_max if config else 7

    # --- Navegación histórica por semanas ---
    try:
        week_offset = int(request.GET.get('week_offset', 0))
    except (TypeError, ValueError):
        week_offset = 0
    # Clamp defensivo: nunca proyectar semanas futuras (>0) ni ir más allá de 12 semanas atrás,
    # sin importar lo que el usuario manipule manualmente en la URL.
    week_offset = max(min(week_offset, 0), -12)

    hoy = date.today()
    lunes_semana_actual = hoy - timedelta(days=hoy.weekday())
    fecha_inicio = lunes_semana_actual + timedelta(weeks=week_offset)
    fecha_fin = fecha_inicio + timedelta(days=13)
    lunes_esta_semana = fecha_inicio

    # 1. Traer excepciones del rango (recalculado sobre la ventana histórica seleccionada)
    excepciones = {
        cap.fecha: cap.limite_personalizado
        for cap in CapacidadDiaria.objects.filter(fecha__range=[fecha_inicio, fecha_fin])
    }

    # 2. Consultar despachos (recalculado sobre la ventana histórica seleccionada)
    despachos = PlanificacionDespacho.objects.filter(
        fecha_despacho__range=[fecha_inicio, fecha_fin]
    ).order_by('fecha_despacho', 'cliente')

    datos_por_dia = collections.defaultdict(list)
    conteo_por_fecha = collections.defaultdict(set) 
    
    for d in despachos:
        datos_por_dia[d.fecha_despacho].append(d)
        conteo_por_fecha[d.fecha_despacho].add(d.unidad_placa)

    agenda = collections.defaultdict(list)
    
    # 3. Construir calendario
    for i in range(14):
        dia_actual = fecha_inicio + timedelta(days=i)
        if dia_actual.weekday() == 6: continue

        registros_dia = datos_por_dia.get(dia_actual, [])
        conteo_placas = len(conteo_por_fecha.get(dia_actual, set()))
        
        # LÓGICA CLAVE: Jerarquía de límites
        limite_dia = excepciones.get(dia_actual, capacidad_limite_global)
        
        # Porcentaje basado en el límite de este día
        porcentaje = (conteo_placas / limite_dia) * 100 if limite_dia > 0 else 0
        
        # SEMÁFORO SINCRONIZADO:
        # Rojo si SUPERA el límite del día. Ámbar si alcanza el 100% o está cerca (>=80%).
        if conteo_placas > limite_dia:
            estado, color_class = "SOBRESATURADO", "rose"
            estilo_card = "border-rose-500 shadow-lg shadow-rose-100/50"
            estilo_header = "bg-rose-600 text-white"
        elif porcentaje >= 100:
            estado, color_class = "CRÍTICO", "amber"
            estilo_card = "border-amber-500 shadow-md shadow-amber-50/50"
            estilo_header = "bg-amber-600 text-white"
        elif porcentaje >= 80:
            estado, color_class = "ADVERTENCIA", "amber"
            estilo_card = "border-amber-300"
            estilo_header = "bg-amber-500 text-white"
        else:
            estado, color_class = "NORMAL", "emerald"
            estilo_card = "border-slate-100"
            estilo_header = "bg-slate-100 text-slate-700"

        # Agrupación de empresas para el detalle del card
        dict_empresas = {}
        for r in registros_dia:
            if r.cliente not in dict_empresas:
                dict_empresas[r.cliente] = {'nombre': r.cliente, 'id_referencia': r.id}
        
        dia_info = {
            'fecha': dia_actual,
            'nombre_dia': obtener_nombre_dia_es(dia_actual),
            'mes': obtener_nombre_mes_es(dia_actual),
            'empresas': sorted(dict_empresas.values(), key=lambda x: x['nombre']),
            'conteo_placas': conteo_placas,
            'limite_dia': limite_dia, 
            'porcentaje': porcentaje,
            'estado': estado,
            'color_class': color_class,
            'estilo_card': estilo_card,
            'estilo_header': estilo_header,
        }
        
        num_semana = dia_actual.isocalendar()[1]
        agenda[num_semana].append(dia_info)
    
    # KPIs con excepciones pasadas por parámetro para exactitud
    kpis_data = calcular_kpis_dashboard(lunes_esta_semana, conteo_por_fecha, capacidad_limite_global, excepciones)

    return render(request, 'planificacion/dashboard.html', {
        'agenda': dict(agenda),
        'capacidad_limite': capacidad_limite_global,
        'kpis': kpis_data,
        # Estado de navegación histórica para los controles << / >> del template
        'week_offset': week_offset,
        'es_semana_actual': week_offset == 0,
        'es_limite_historico': week_offset <= -12,
        'week_offset_prev': max(week_offset - 1, -12),
        'week_offset_next': min(week_offset + 1, 0),
        'rango_fecha_inicio': fecha_inicio,
        'rango_fecha_fin': fecha_fin,
    })

@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def actualizar_capacidad(request):
    """Ajusta el límite máximo de la flota global."""
    if request.method == 'POST':
        nuevo_limite = request.POST.get('nuevo_limite')
        if nuevo_limite:
            config, created = ConfiguracionCapacidad.objects.get_or_create(id=1)
            config.limite_max = int(nuevo_limite)
            config.save()
            messages.success(request, f"Capacidad de flota actualizada a {nuevo_limite} unidades.")
    return redirect('dashboard')

@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def gestionar_capacidad_diaria(request):
    """
    Vista rápida para que Planeamiento ajuste el límite de camiones por fecha.
    Recibe 'fecha' (YYYY-MM-DD) y 'limite_personalizado' (int) vía POST.
    """
    if request.method == 'POST':
        fecha_str = request.POST.get('fecha')
        nuevo_limite = request.POST.get('limite_personalizado')
        motivo = request.POST.get('motivo', 'Ajuste manual desde Dashboard')

        if not fecha_str or not nuevo_limite:
            messages.error(request, "⚠️ Datos incompletos para actualizar la capacidad.")
            return redirect('dashboard')

        try:
            # 1. Validación de formato y valor numérico
            limite_int = int(nuevo_limite)
            if limite_int < 0:
                messages.error(request, "⚠️ El límite no puede ser negativo.")
                return redirect('dashboard')

            # 2. Validación de Integridad Histórica (Opcional pero Recomendado)
            # Evita que se cambie la capacidad de un día que ya pasó
            from datetime import datetime
            fecha_dt = datetime.strptime(fecha_str, '%Y-%m-%d').date()
            if fecha_dt < timezone.now().date():
                messages.error(request, "🚫 No se puede modificar la capacidad de fechas pasadas.")
                return redirect('dashboard')

            # 3. Lógica central: Actualiza si existe, crea si no.
            obj, created = CapacidadDiaria.objects.update_or_create(
                fecha=fecha_dt,
                defaults={
                    'limite_personalizado': limite_int,
                    'motivo': motivo
                }
            )

            accion = "establecida" if created else "actualizada"
            messages.success(request, f"✅ Capacidad para el {fecha_str} {accion} a {limite_int} unidades.")

        except ValueError:
            messages.error(request, "⚠️ El límite debe ser un número entero válido.")
        except Exception as e:
            messages.error(request, f"❌ Error al procesar el ajuste: {str(e)}")

    return redirect('dashboard')


@login_required
def detalle_dia(request, fecha):
    """Muestra todos los despachos de una fecha específica con su límite real."""
    fecha_dt = date.fromisoformat(fecha)
    despachos = PlanificacionDespacho.objects.filter(fecha_despacho=fecha_dt).order_by('cliente')
    
    # Obtener límite real para este día
    excepcion = CapacidadDiaria.objects.filter(fecha=fecha_dt).first()
    config = ConfiguracionCapacidad.objects.first()
    limite_dia = excepcion.limite_personalizado if excepcion else (config.limite_max if config else 7)
    
    placas_unicas = set(d.unidad_placa for d in despachos)
    conteo_placas = len(placas_unicas)

    # Totales del día (registros y pallets) vía agregación ORM
    totales = despachos.aggregate(total_registros=Count('id'), total_pallets=Sum('pallets'))
    total_registros = totales['total_registros']
    total_pallets = totales['total_pallets'] or 0

    es_editable = fecha_dt >= date.today()

    return render(request, 'planificacion/detalle_dia.html', {
        'despachos': despachos,
        'fecha': fecha_dt,
        'conteo_placas': conteo_placas,
        'capacidad_limite': limite_dia, # Pasamos el límite real de ese día
        'nombre_dia': obtener_nombre_dia_es(fecha_dt),
        'es_editable': es_editable,
        'usuario_puede_editar': es_planeamiento(request.user),
        'total_registros': total_registros,
        'total_pallets': total_pallets,
    })


@login_required
def exportar_detalle_dia_excel(request, fecha):
    """Exporta a Excel los despachos de una fecha específica."""
    fecha_dt = date.fromisoformat(fecha)
    despachos = PlanificacionDespacho.objects.filter(fecha_despacho=fecha_dt).order_by('cliente')
    nombre_archivo = f"despachos_{fecha_dt.isoformat()}.xlsx"
    return _generar_excel_despachos(despachos, nombre_archivo)

def _construir_queryset_busqueda(request):
    """
    Construye el queryset de PlanificacionDespacho filtrado según los
    parámetros GET de la request (cliente, placa, fecha_desde, fecha_hasta),
    aplicando el filtro de fechas por defecto (mes actual) cuando corresponde.
    Compartido entre la vista paginada "Consultar" y su exportación a Excel,
    para que ambas apliquen siempre exactamente los mismos filtros.
    """
    queryset = PlanificacionDespacho.objects.all().order_by('-fecha_despacho')

    hoy = date.today()
    primer_dia_mes = hoy.replace(day=1)

    q_cliente = request.GET.get('cliente', '')
    q_placa = request.GET.get('placa', '')

    # Filtro de fechas por defecto: solo se aplica si NINGUNA de las dos claves
    # llegó en request.GET (primera visita / enlace "Limpiar Filtros"). Si el
    # usuario envía explícitamente los campos vacíos (ej. buscando solo por
    # cliente), se respeta el filtro abierto tal como llegó, sin fecha límite.
    if 'fecha_desde' in request.GET or 'fecha_hasta' in request.GET:
        q_desde = request.GET.get('fecha_desde', '')
        q_hasta = request.GET.get('fecha_hasta', '')
    else:
        q_desde = primer_dia_mes.isoformat()
        q_hasta = hoy.isoformat()

    if q_cliente:
        queryset = queryset.filter(cliente__icontains=q_cliente)
    if q_placa:
        queryset = queryset.filter(unidad_placa__icontains=q_placa)
    if q_desde:
        queryset = queryset.filter(fecha_despacho__gte=q_desde)
    if q_hasta:
        queryset = queryset.filter(fecha_despacho__lte=q_hasta)

    filtros = {
        'cliente': q_cliente,
        'placa': q_placa,
        'fecha_desde': q_desde,
        'fecha_hasta': q_hasta,
    }
    return queryset, filtros


@login_required
def lista_busqueda(request):
    """Vista de gestión y búsqueda con filtros, paginada (10 registros por página)."""
    queryset, filtros = _construir_queryset_busqueda(request)

    # Totales sobre el queryset COMPLETO ya filtrado, antes de paginar
    # (una sola consulta de agregación, no derivada de la página actual).
    totales = queryset.aggregate(total_registros=Count('id'), total_pallets=Sum('pallets'))
    total_registros = totales['total_registros']
    total_pallets = totales['total_pallets'] or 0

    # Querystring de filtros sin 'page', para no duplicar el parámetro
    # en los enlaces de paginación (?page=N&cliente=...).
    querydict = request.GET.copy()
    querydict.pop('page', None)
    querystring = querydict.urlencode()

    paginator = Paginator(queryset, REGISTROS_POR_PAGINA)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'planificacion/lista_busqueda.html', {
        'page_obj': page_obj,
        'querystring': querystring,
        'total_registros': total_registros,
        'total_pallets': total_pallets,
        'filtros': filtros,
        'usuario_puede_editar': es_planeamiento(request.user),
    })


@login_required
def exportar_busqueda_excel(request):
    """Exporta a Excel el queryset COMPLETO filtrado de 'Consultar' (no solo la página actual)."""
    queryset, filtros = _construir_queryset_busqueda(request)
    nombre_archivo = f"despachos_{filtros['fecha_desde'] or 'inicio'}_a_{filtros['fecha_hasta'] or 'hoy'}.xlsx"
    return _generar_excel_despachos(queryset, nombre_archivo)

@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def registrar_planificacion(request):
    """Registro de nuevas planificaciones con validación de capacidad dinámica."""
    fecha_predefinida = request.GET.get('fecha')
    
    if request.method == 'POST':
        form = PlanificacionForm(request.POST)
        if form.is_valid():
            nueva_fecha = form.cleaned_data.get('fecha_despacho')
            if nueva_fecha < date.today():
                messages.error(request, "⚠️ Error de Seguridad: No se pueden crear registros con fechas pasadas.")
                return render(request, 'planificacion/form_planificacion.html', {'form': form})

            planificacion = form.save(commit=False)
            planificacion.creado_por = request.user
            planificacion.save()

            # Lógica de alerta basada en Capacidad Diaria
            excepcion = CapacidadDiaria.objects.filter(fecha=nueva_fecha).first()
            config = ConfiguracionCapacidad.objects.first()
            limite_real = excepcion.limite_personalizado if excepcion else (config.limite_max if config else 7)
            
            total_placas = PlanificacionDespacho.objects.filter(fecha_despacho=nueva_fecha).values('unidad_placa').distinct().count()
            
            if total_placas > limite_real:
                messages.warning(request, f"⚠️ Unidad {planificacion.unidad_placa} registrada, pero se ha superado el límite para este día ({total_placas}/{limite_real}).")
            else:
                messages.success(request, f"✅ Unidad {planificacion.unidad_placa} programada correctamente.")

            return redirect('dashboard')
        else:
            messages.error(request, "Error en el formulario. Verifique los datos.")
    else:
        initial_data = {}
        if fecha_predefinida:
            initial_data['fecha_despacho'] = fecha_predefinida
        form = PlanificacionForm(initial=initial_data)

    return render(request, 'planificacion/form_planificacion.html', {'form': form, 'edit_mode': False})

@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def editar_planificacion(request, pk):
    """Edición de planificaciones existentes con validación de capacidad dinámica."""
    registro = get_object_or_404(PlanificacionDespacho, pk=pk)
    
    if registro.fecha_despacho < date.today():
        messages.error(request, "⚠️ No se pueden editar registros de fechas pasadas.")
        return redirect('dashboard')

    if request.method == 'POST':
        form = PlanificacionForm(request.POST, instance=registro)
        if form.is_valid():
            form.save()
            
            # Obtener límite real del día editado
            excepcion = CapacidadDiaria.objects.filter(fecha=registro.fecha_despacho).first()
            config = ConfiguracionCapacidad.objects.first()
            limite_real = excepcion.limite_personalizado if excepcion else (config.limite_max if config else 7)
            
            total_placas = PlanificacionDespacho.objects.filter(fecha_despacho=registro.fecha_despacho).values('unidad_placa').distinct().count()
            
            if total_placas > limite_real:
                messages.warning(request, f"⚠️ Planificación actualizada. Nota: La flota actual ({total_placas}) excede el límite de este día ({limite_real}).")
            else:
                messages.success(request, "✅ Planificación actualizada correctamente.")
            return redirect('dashboard')
    else:
        form = PlanificacionForm(instance=registro)

    return render(request, 'planificacion/form_planificacion.html', {'form': form, 'edit_mode': True})

@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def eliminar_planificacion(request, pk):
    """Eliminación de registros con bloqueo histórico."""
    registro = get_object_or_404(PlanificacionDespacho, pk=pk)
    
    if registro.fecha_despacho < date.today():
        messages.error(request, "🚫 No se pueden eliminar registros de fechas pasadas.")
        return redirect('lista_busqueda')
        
    if request.method == 'POST':
        placa = registro.unidad_placa
        fecha = registro.fecha_despacho
        registro.delete()
        messages.success(request, f"🗑️ Se eliminó la programación de la unidad {placa} para el {fecha}.")

        referer = request.META.get('HTTP_REFERER')
        if referer and 'busqueda' in referer:
            return redirect('lista_busqueda')
        return redirect('dashboard')

    return redirect('dashboard')


# --- MÓDULO DE CARGA MASIVA POR EXCEL ---

@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def carga_masiva_excel(request):
    """
    Carga masiva de despachos desde un archivo .xlsx, procesado en memoria
    (BytesIO). Reutiliza PlanificacionForm por cada fila para garantizar
    exactamente las mismas validaciones de negocio que el alta manual
    (placa duplicada el mismo día, límites de campo, etc.), además de un
    chequeo explícito de fecha futura con detalle de error por fila.
    """
    resultado = None

    if request.method == 'POST':
        archivo = request.FILES.get('archivo_excel')

        if not archivo:
            messages.error(request, "⚠️ Debes seleccionar un archivo .xlsx para cargar.")
        elif not archivo.name.lower().endswith('.xlsx'):
            messages.error(request, "⚠️ El archivo debe tener extensión .xlsx (Excel moderno).")
        else:
            try:
                wb = load_workbook(filename=BytesIO(archivo.read()), data_only=True)
                ws = wb.active
            except Exception:
                messages.error(request, "❌ No se pudo leer el archivo. Verifica que sea un Excel (.xlsx) válido y no esté dañado.")
                ws = None

            if ws is not None:
                # 1. Validar encabezados por NOMBRE (fila 1), no por posición fija de columna
                primera_fila = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                encabezados_archivo = [str(c).strip().lower() if c else '' for c in primera_fila]
                columnas_faltantes = [c for c in COLUMNAS_EXCEL_CARGA_MASIVA if c not in encabezados_archivo]

                if columnas_faltantes:
                    messages.error(
                        request,
                        f"❌ Faltan columnas requeridas en el archivo: {', '.join(columnas_faltantes)}. "
                        f"Descarga la plantilla para ver la estructura exacta."
                    )
                else:
                    indice_columna = {nombre: encabezados_archivo.index(nombre) for nombre in COLUMNAS_EXCEL_CARGA_MASIVA}
                    total_filas = max((ws.max_row or 1) - 1, 0)

                    if total_filas > MAX_FILAS_CARGA_MASIVA:
                        messages.error(
                            request,
                            f"⚠️ El archivo tiene {total_filas} filas de datos; el máximo permitido por carga "
                            f"es {MAX_FILAS_CARGA_MASIVA}. Divide el archivo en lotes más pequeños."
                        )
                    else:
                        exitosos = 0
                        errores = []

                        for row_num, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                            # Saltar filas completamente vacías (comunes al final de una plantilla)
                            if not fila or all(c is None or str(c).strip() == '' for c in fila):
                                continue

                            valores = {
                                nombre: (fila[idx] if idx < len(fila) else None)
                                for nombre, idx in indice_columna.items()
                            }

                            # 2. Parseo estricto de fecha AAAAMMDD -> date
                            try:
                                fecha_parseada = _parsear_fecha_aaaammdd(valores['fecha'])
                            except ValueError as e:
                                errores.append(f"Fila {row_num}: {e}. Se espera formato AAAAMMDD, ej: 20260820.")
                                continue

                            # 3. Regla de negocio explícita: solo fechas de hoy en adelante
                            if fecha_parseada < date.today():
                                errores.append(
                                    f"Fila {row_num}: la fecha {fecha_parseada.isoformat()} ya pasó. "
                                    f"Solo se permiten despachos de hoy en adelante."
                                )
                                continue

                            # 4. Resto de validaciones reutilizando PlanificacionForm
                            #    (placa duplicada el mismo día, longitudes de campo, etc.)
                            form_data = {
                                'fecha_despacho': fecha_parseada,
                                'cliente': valores['cliente'],
                                'sede': valores['sede_destino'],
                                'unidad_placa': valores['placa'],
                                'conductor': valores['conductor'],
                                'orden_compra': valores['orden_compra'],
                                'carga_origen': valores['origen'],
                                'pallets': valores['pallets'],
                                'costo_transporte': valores['costo_transporte'] or 0,
                            }

                            form = PlanificacionForm(form_data)
                            if form.is_valid():
                                despacho = form.save(commit=False)
                                despacho.creado_por = request.user
                                despacho.save()
                                exitosos += 1
                            else:
                                detalle = '; '.join(
                                    f"{campo}: {' / '.join(errs)}" for campo, errs in form.errors.items()
                                )
                                errores.append(f"Fila {row_num}: {detalle}")

                        resultado = {
                            'total_filas': total_filas,
                            'exitosos': exitosos,
                            'fallidos': len(errores),
                            'errores': errores,
                        }

                        if exitosos and not errores:
                            messages.success(request, f"✅ Se cargaron correctamente {exitosos} despachos.")
                        elif exitosos:
                            messages.warning(request, f"⚠️ Se cargaron {exitosos} de {total_filas} filas. Revisa el detalle de errores.")
                        else:
                            messages.error(request, "❌ No se pudo cargar ninguna fila. Revisa el detalle de errores.")

    return render(request, 'planificacion/carga_masiva.html', {
        'columnas_esperadas': COLUMNAS_EXCEL_CARGA_MASIVA,
        'resultado': resultado,
    })


@login_required
@user_passes_test(es_planeamiento, login_url='dashboard')
def descargar_plantilla_excel(request):
    """Genera y descarga una plantilla .xlsx vacía (con fila de ejemplo) para la carga masiva."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Plantilla'

    relleno_header = PatternFill(start_color=COLOR_HEADER_EXCEL, end_color=COLOR_HEADER_EXCEL, fill_type='solid')
    fuente_header = Font(color='FFFFFF', bold=True)
    for col_num, columna in enumerate(COLUMNAS_EXCEL_CARGA_MASIVA, start=1):
        celda = ws.cell(row=1, column=col_num, value=columna)
        celda.fill = relleno_header
        celda.font = fuente_header
        ws.column_dimensions[get_column_letter(col_num)].width = max(len(columna) + 4, 16)

    fila_ejemplo = [
        date.today().strftime('%Y%m%d'), 'ABC-123', 'Juan Pérez', 'Cliente Ejemplo',
        'Sede Ejemplo', 'Almacén Central', 'OC-0001', 10, 150.00,
    ]
    for col_num, valor in enumerate(fila_ejemplo, start=1):
        ws.cell(row=2, column=col_num, value=valor)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="plantilla_carga_masiva.xlsx"'
    return response


# --- MÓDULO DE ESTADÍSTICA (SOLO LECTURA) ---

@login_required
def estadisticas_dashboard(request):
    """
    Módulo de solo lectura "Estadística": reportes y paneles analíticos
    (Dashboards Pro) filtrables por rango de fechas y cliente.
    No realiza ninguna operación de escritura sobre la base de datos.
    """
    hoy = date.today()

    # --- 1. Filtros GET, con manejo seguro de valores vacíos/inválidos ---
    cliente_filtro = (request.GET.get('cliente') or '').strip()

    fecha_desde_raw = request.GET.get('fecha_desde')
    try:
        fecha_desde = date.fromisoformat(fecha_desde_raw) if fecha_desde_raw else hoy - timedelta(days=30)
    except ValueError:
        fecha_desde = hoy - timedelta(days=30)

    fecha_hasta_raw = request.GET.get('fecha_hasta')
    try:
        fecha_hasta = date.fromisoformat(fecha_hasta_raw) if fecha_hasta_raw else hoy
    except ValueError:
        fecha_hasta = hoy

    # Defensa: si el usuario invierte el rango en la URL, lo corregimos en vez de fallar.
    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde

    # Clamp defensivo de rendimiento: limita el recorrido día a día a un máximo razonable.
    rango_truncado = False
    if (fecha_hasta - fecha_desde).days > MAX_RANGO_ESTADISTICAS_DIAS:
        fecha_desde = fecha_hasta - timedelta(days=MAX_RANGO_ESTADISTICAS_DIAS)
        rango_truncado = True

    # --- 2. Queryset base filtrado, reutilizado por todos los paneles ---
    despachos = PlanificacionDespacho.objects.filter(
        fecha_despacho__range=[fecha_desde, fecha_hasta]
    )
    if cliente_filtro:
        despachos = despachos.filter(cliente__icontains=cliente_filtro)

    # ================================================================
    # PANEL 1: Volumen y Tendencia de Ocupación Histórica
    # Placas únicas operadas por día (ORM) vs. límite dinámico de capacidad.
    # ================================================================
    config = ConfiguracionCapacidad.objects.first()
    capacidad_limite_global = config.limite_max if config else 7

    capacidades_dia = list(CapacidadDiaria.objects.filter(fecha__range=[fecha_desde, fecha_hasta]))
    excepciones_rango = {c.fecha: c.limite_personalizado for c in capacidades_dia}

    conteo_diario = (
        despachos.values('fecha_despacho')
        .annotate(unidades=Count('unidad_placa', distinct=True))
    )
    unidades_por_fecha = {c['fecha_despacho']: c['unidades'] for c in conteo_diario}

    panel1_labels, panel1_ocupacion, panel1_limite = [], [], []
    dias_criticos_rango = 0
    dia_cursor = fecha_desde
    while dia_cursor <= fecha_hasta:
        unidades = unidades_por_fecha.get(dia_cursor, 0)
        limite_dia = excepciones_rango.get(dia_cursor, capacidad_limite_global)

        panel1_labels.append(dia_cursor.strftime('%d/%m'))
        panel1_ocupacion.append(unidades)
        panel1_limite.append(limite_dia)

        if limite_dia > 0 and unidades >= limite_dia:
            dias_criticos_rango += 1

        dia_cursor += timedelta(days=1)

    # ================================================================
    # PANELES 2 y 4: se apoyan en un único fetch de registros crudos para
    # evitar múltiples idas y vueltas a la base de datos sobre el mismo
    # queryset filtrado (el pivote y la "moda" no son operaciones nativas del ORM).
    # ================================================================
    registros = list(despachos.values('cliente', 'unidad_placa', 'conductor', 'fecha_despacho', 'pallets'))

    # --- PANEL 2: Matriz de Concentración (Cliente x Día de la Semana) ---
    matriz_raw = collections.defaultdict(lambda: collections.defaultdict(int))
    for r in registros:
        dia_nombre = DIAS_SEMANA_ES[r['fecha_despacho'].weekday()]
        matriz_raw[r['cliente']][dia_nombre] += r['pallets']

    matriz_dias_semana = DIAS_SEMANA_ES[:6]  # Lunes a Sábado (Domingo sin operación, igual que el Dashboard)
    matriz_cliente_dia = []
    for cliente, valores_dia in matriz_raw.items():
        fila = {
            'cliente': cliente,
            'valores': [valores_dia.get(dia, 0) for dia in matriz_dias_semana],
        }
        fila['total'] = sum(fila['valores'])
        matriz_cliente_dia.append(fila)
    matriz_cliente_dia.sort(key=lambda f: f['total'], reverse=True)

    # --- PANEL 4: Auditoría y Frecuencia de Uso de Flota ---
    flota_raw = collections.defaultdict(lambda: {
        'dias': set(), 'servicios': 0, 'clientes': set(), 'conductores': collections.Counter()
    })
    for r in registros:
        info = flota_raw[r['unidad_placa']]
        info['dias'].add(r['fecha_despacho'])
        info['servicios'] += 1
        info['clientes'].add(r['cliente'])
        info['conductores'][r['conductor']] += 1

    rendimiento_flota = []
    for placa, info in flota_raw.items():
        conductor_frecuente = info['conductores'].most_common(1)[0][0] if info['conductores'] else '-'
        rendimiento_flota.append({
            'placa': placa,
            'conductor_frecuente': conductor_frecuente,
            'dias_trabajados': len(info['dias']),
            'servicios_totales': info['servicios'],
            'clientes_atendidos': len(info['clientes']),
        })
    rendimiento_flota.sort(key=lambda f: f['servicios_totales'], reverse=True)

    # ================================================================
    # PANEL 3: Reporte de Eficiencia de Carga y Densidad de Pallets
    # 100% agregaciones ORM (Sum, Avg, Count).
    # ================================================================
    resumen_pallets = despachos.aggregate(
        total_pallets=Sum('pallets'),
        promedio_pallets=Avg('pallets'),
        total_despachos=Count('id'),
    )
    resumen_pallets['total_pallets'] = resumen_pallets['total_pallets'] or 0
    resumen_pallets['promedio_pallets'] = (
        round(resumen_pallets['promedio_pallets'], 1) if resumen_pallets['promedio_pallets'] else 0
    )

    top_origenes_qs = (
        despachos.exclude(carga_origen='')
        .values('carga_origen')
        .annotate(total_pallets=Sum('pallets'))
        .order_by('-total_pallets')[:8]
    )
    top_origenes_labels = [o['carga_origen'] for o in top_origenes_qs]
    top_origenes_valores = [o['total_pallets'] for o in top_origenes_qs]

    # ================================================================
    # PANEL 5: Histórico de Saturación e Incidencias de Capacidad
    # ================================================================
    motivos_qs = (
        CapacidadDiaria.objects.filter(fecha__range=[fecha_desde, fecha_hasta])
        .exclude(motivo__isnull=True).exclude(motivo__exact='')
        .values('motivo')
        .annotate(total=Count('id'))
        .order_by('-total')
    )
    motivos_labels = [m['motivo'] for m in motivos_qs]
    motivos_valores = [m['total'] for m in motivos_qs]

    context = {
        'filtros': {
            'fecha_desde': fecha_desde,
            'fecha_hasta': fecha_hasta,
            'cliente': cliente_filtro,
        },
        'rango_truncado': rango_truncado,
        'max_rango_dias': MAX_RANGO_ESTADISTICAS_DIAS,
        # Panel 1
        'panel1_labels': panel1_labels,
        'panel1_ocupacion': panel1_ocupacion,
        'panel1_limite': panel1_limite,
        'dias_criticos_rango': dias_criticos_rango,
        # Panel 2
        'matriz_dias_semana': matriz_dias_semana,
        'matriz_cliente_dia': matriz_cliente_dia,
        # Panel 3
        'resumen_pallets': resumen_pallets,
        'top_origenes_labels': top_origenes_labels,
        'top_origenes_valores': top_origenes_valores,
        # Panel 4
        'rendimiento_flota': rendimiento_flota,
        # Panel 5
        'motivos_labels': motivos_labels,
        'motivos_valores': motivos_valores,
    }
    return render(request, 'planificacion/estadisticas.html', context)
