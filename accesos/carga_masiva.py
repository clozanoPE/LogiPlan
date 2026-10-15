"""
Carga masiva por Excel de Transportes, Conductores, Programaciones y licencias de conductores.

Mismo principio que la carga masiva de Planificación: cada fila se valida con el MISMO
formulario del alta manual (TransporteForm, ConductorForm, ProgramacionForm), así que rigen
exactamente las mismas reglas. Diferencias propias de este módulo:
- Dos pasos: primero una vista previa (no graba nada) y luego la confirmación, que vuelve a
  validar cada fila contra el estado actual de la base y graba solo las válidas.
- En los maestros, un registro que ya existe (placa o documento) se omite y se informa; no se
  actualiza. La única carga que modifica registros existentes es la de licencias (CargaLicencias),
  que solo toca la licencia y las notificaciones de conductores ya registrados.
- Las columnas se reconocen por nombre (sin importar mayúsculas, tildes ni espacios) y las
  celdas que Excel convierte (fechas, horas, números) se normalizan antes de validar.
Todo se procesa en memoria; no se guarda el archivo.
"""
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import auditoria, jornadas, licencias, servicios, tiempo
from .forms import ConductorForm, ProgramacionForm, TransporteForm
from .models import (CATEGORIAS_LICENCIA, CATEGORIAS_LICENCIA_ESPECIAL, CentroDistribucion, Conductor,
                     ConfiguracionAccesos, Programacion, TamanoVehiculo, TipoDocumentoIdentidad, Transporte,
                     normalizar_placa)

MAX_FILAS = 500
# Excel quita los ceros iniciales de un documento escrito como número; solo se reponen hasta dos
# (un DNI 08123456 llega como 8123456). Un número más corto es un error de digitación.
CEROS_REPONIBLES = 2
COLOR_CABECERA = 'EA580C'  # naranja del módulo
FILAS_CON_FORMATO = MAX_FILAS + 1

OK = 'OK'
AVISO = 'AVISO'
OMITIDA = 'OMITIDA'
ERROR = 'ERROR'


class ArchivoInvalido(Exception):
    """El archivo no se puede procesar (formato, columnas, tamaño)."""


@dataclass
class Columna:
    nombre: str
    titulo: str
    obligatoria: bool
    formato: str          # texto de ayuda para la hoja de instrucciones
    ejemplo: object
    tipo: str = 'texto'   # texto | fecha | hora | si_no | documento
    lista: str = ''       # nombre de la lista de valores válidos (hoja «Listas»)


@dataclass
class Fila:
    numero: int
    valores: dict
    estado: str = OK
    mensajes: list = field(default_factory=list)
    form: object = None

    def error(self, mensaje):
        self.estado = ERROR
        self.mensajes.append(mensaje)
        return self

    def omitir(self, mensaje):
        self.estado = OMITIDA
        self.mensajes.append(mensaje)
        return self

    @property
    def grabable(self):
        return self.estado in (OK, AVISO)

    def resumen(self):
        return {'numero': self.numero, 'estado': self.estado, 'mensajes': self.mensajes,
                'valores': self.valores}


# --- NORMALIZACIÓN DE CELDAS ---

def clave(texto):
    """'Razón Social ' -> 'razon_social'; también para comparar valores de listas sin tildes."""
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return '_'.join(texto.strip().lower().replace('-', ' ').replace('/', ' ').split())


def a_texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d %H:%M')
    if isinstance(valor, (date, time)):
        return valor.isoformat()
    return str(valor).strip()


def a_fecha(valor):
    """Acepta fecha de Excel, DD/MM/AAAA, AAAA-MM-DD, AAAAMMDD y el número de serie de Excel."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        numero = int(valor)
        if 19000101 <= numero <= 29991231:
            valor = str(numero)
        elif 20000 <= numero <= 80000:  # número de serie de Excel (días desde 1899-12-30)
            return date(1899, 12, 30) + timedelta(days=numero)
    texto = a_texto(valor)
    for formato in ('%d/%m/%Y', '%Y-%m-%d', '%Y%m%d', '%d-%m-%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError(f"fecha «{texto}» no válida; use DD/MM/AAAA")


def a_hora(valor):
    """Acepta hora de Excel, fracción de día, HH:MM y HH:MM:SS."""
    if isinstance(valor, datetime):
        return valor.time().replace(second=0, microsecond=0)
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    if isinstance(valor, timedelta):
        minutos = round(valor.total_seconds() / 60)
        return time(minutos // 60 % 24, minutos % 60)
    if isinstance(valor, float) and 0 <= valor < 1:
        minutos = round(valor * 24 * 60)
        return time(minutos // 60 % 24, minutos % 60)
    texto = a_texto(valor).replace('.', ':').replace('h', ':').rstrip(':')
    for formato in ('%H:%M', '%H:%M:%S', '%I:%M %p'):
        try:
            return datetime.strptime(texto.upper(), formato).time().replace(second=0)
        except ValueError:
            continue
    raise ValueError(f"hora «{texto}» no válida; use HH:MM (24 horas)")


def a_si_no(valor):
    texto = clave(a_texto(valor))
    if texto in ('', 'no', 'n', '0', 'false', 'falso'):
        return False
    if texto in ('si', 's', '1', 'true', 'verdadero', 'x'):
        return True
    raise ValueError(f"«{a_texto(valor)}» no es SI ni NO")


def a_documento(valor):
    """Documento como texto, sin espacios ni puntos (Excel lo puede traer como número)."""
    return a_texto(valor).replace(' ', '').replace('.', '').replace('-', '').upper()


CONVERSORES = {'texto': a_texto, 'fecha': a_fecha, 'hora': a_hora, 'si_no': a_si_no, 'documento': a_documento}


def _serializable(valor):
    if isinstance(valor, (date, time)):
        return valor.isoformat()[:5] if isinstance(valor, time) else valor.isoformat()
    return valor


# --- LECTURA DEL ARCHIVO ---

def leer_archivo(archivo, columnas):
    """Devuelve [(número de fila, {columna: valor crudo})]. Lanza ArchivoInvalido si no se puede usar."""
    nombre = getattr(archivo, 'name', '') or ''
    if not nombre.lower().endswith('.xlsx'):
        raise ArchivoInvalido("El archivo debe ser un Excel .xlsx (use la plantilla).")
    try:
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ArchivoInvalido("No se pudo leer el archivo. Verifique que sea un .xlsx válido.") from None
    hoja = libro['Datos'] if 'Datos' in libro.sheetnames else libro.active
    filas = hoja.iter_rows(values_only=True)
    encabezados = [clave(c) for c in next(filas, ())]
    faltantes = [c.titulo for c in columnas if c.obligatoria and c.nombre not in encabezados]
    if faltantes:
        raise ArchivoInvalido("Faltan columnas en el archivo: " + ', '.join(faltantes)
                              + ". Descargue la plantilla y no cambie los encabezados.")
    indices = {c.nombre: encabezados.index(c.nombre) for c in columnas if c.nombre in encabezados}
    leidas = []
    for numero, fila in enumerate(filas, start=2):
        crudos = {nombre: (fila[i] if i < len(fila) else None) for nombre, i in indices.items()}
        if all(a_texto(v) == '' for v in crudos.values()):
            continue  # filas vacías (comunes al final de la plantilla)
        leidas.append((numero, crudos))
        if len(leidas) > MAX_FILAS:
            libro.close()
            raise ArchivoInvalido(f"El archivo tiene más de {MAX_FILAS} filas con datos. Divídalo en partes.")
    libro.close()
    if not leidas:
        raise ArchivoInvalido("El archivo no tiene filas con datos.")
    return leidas


def normalizar(crudos, columnas):
    """Convierte las celdas según el tipo de cada columna. Devuelve (valores, errores)."""
    valores, errores = {}, []
    for col in columnas:
        crudo = crudos.get(col.nombre)
        if a_texto(crudo) == '':
            valores[col.nombre] = '' if col.tipo != 'si_no' else False
            if col.obligatoria:
                errores.append(f"falta {col.titulo}")
            continue
        try:
            valores[col.nombre] = _serializable(CONVERSORES[col.tipo](crudo))
        except ValueError as error:
            valores[col.nombre] = a_texto(crudo)
            errores.append(f"{col.titulo}: {error}")
    return valores, errores


def _errores_del_form(form):
    mensajes = []
    for campo, errores in form.errors.items():
        etiqueta = form.fields[campo].label if campo in form.fields else ''
        for error in errores:
            mensajes.append(f"{etiqueta}: {error}" if etiqueta else error)
    return mensajes


# --- TIPOS DE CARGA ---

class TipoCarga:
    clave = ''
    titulo = ''
    singular = ''
    columnas = []
    nota = ''
    # Textos de pantalla (la carga de licencias actualiza en vez de registrar).
    participio = 'Registrada'
    verbo_plural = 'registraron'
    volver_a = ''
    titulo_propio = ''

    @property
    def titulo_pantalla(self):
        return self.titulo_propio or f"Carga masiva de {self.titulo.lower()}"

    @property
    def destino_titulo(self):
        return self.volver_a or self.titulo

    def listas(self):
        """{nombre de lista: [valores válidos]} para la hoja «Listas» y los desplegables."""
        return {}

    def preparar(self):
        """Datos de consulta compartidos por todas las filas (catálogos, etc.)."""
        return {}

    def validar(self, fila, contexto, vistos, ahora):
        raise NotImplementedError

    def grabar(self, fila, usuario, motivo):
        objeto = fila.form.save(commit=False)
        objeto.creado_por = usuario
        objeto.save()
        auditoria.registrar_creacion(usuario, objeto, motivo=motivo)
        return objeto

    # -- proceso completo --

    def _revisar(self, crudas, ahora):
        contexto = self.preparar()
        vistos = {}
        for numero, crudos in crudas:
            valores, errores = normalizar(crudos, self.columnas)
            fila = Fila(numero, valores)
            if errores:
                fila.estado = ERROR
                fila.mensajes = [e[0].upper() + e[1:] for e in errores]
            else:
                self.validar(fila, contexto, vistos, ahora)
            yield fila

    def evaluar(self, crudas, ahora):
        """Vista previa: valida todas las filas sin grabar nada."""
        return list(self._revisar(crudas, ahora))

    def confirmar(self, crudas, usuario, ahora, nombre_archivo):
        """Vuelve a validar cada fila contra el estado actual (puede haber cambiado desde la vista
        previa) y graba solo las válidas; cada una en su propia transacción."""
        motivo = f"Carga masiva: {nombre_archivo}"[:200]
        filas = []
        # Se valida y graba fila por fila: así las filas siguientes ven las ya grabadas (duplicados, avisos).
        for fila in self._revisar(crudas, ahora):
            if fila.grabable:
                try:
                    with transaction.atomic():
                        objeto = self.grabar(fila, usuario, motivo)
                except IntegrityError:
                    fila.error("Otro usuario registró el mismo dato mientras se cargaba el archivo.")
                else:
                    fila.objeto = objeto
                    self.despues_de_grabar(fila, objeto)
            filas.append(fila)
        return filas

    def despues_de_grabar(self, fila, objeto):
        pass

    # -- plantilla --

    def ejemplo(self):
        return [c.ejemplo for c in self.columnas]


def _buscar(catalogo, texto):
    """Busca en {clave normalizada: objeto} aceptando código o nombre, sin tildes ni mayúsculas."""
    return catalogo.get(clave(texto))


class CargaTransportes(TipoCarga):
    clave = 'transportes'
    titulo = 'Transportes'
    singular = 'transporte'
    nota = "Las placas que ya existen en el maestro se omiten (no se modifican)."
    columnas = [
        Columna('placa', 'placa', True, "Placa del vehículo, con o sin guion (5 a 8 letras o números).", 'F7J-011'),
        Columna('tamano', 'tamano', True, "Tamaño del vehículo: uno de la lista.", 'Mediano', lista='tamanos'),
        Columna('interno', 'interno', False, "SI si es transporte de Daryza; NO o vacío si es de un tercero.",
                'NO', tipo='si_no', lista='si_no'),
    ]

    def listas(self):
        return {'tamanos': list(TamanoVehiculo.objects.filter(activo=True).values_list('nombre', flat=True)),
                'si_no': ['SI', 'NO']}

    def preparar(self):
        return {'tamanos': {clave(t.nombre): t for t in TamanoVehiculo.objects.filter(activo=True)},
                'existentes': set(Transporte.objects.values_list('placa_normalizada', flat=True))}

    def validar(self, fila, contexto, vistos, ahora):
        v = fila.valores
        placa = normalizar_placa(v['placa'])
        if placa in contexto['existentes']:
            return fila.omitir(f"La placa {v['placa']} ya está registrada; se omite.")
        if placa in vistos:
            return fila.omitir(f"Placa repetida en el archivo (fila {vistos[placa]}); se omite.")
        if not contexto['tamanos']:
            return fila.error("No hay tamaños de vehículo registrados. El Administrador debe crearlos en "
                              "Catálogos (Django Admin) antes de cargar transportes.")
        tamano = _buscar(contexto['tamanos'], v['tamano'])
        if tamano is None:
            return fila.error(f"Tamaño «{v['tamano']}» no existe. Use: "
                              + ', '.join(t.nombre for t in contexto['tamanos'].values()) + '.')
        fila.form = TransporteForm({'placa': v['placa'], 'tamano': tamano.pk, 'es_interno': v['interno'],
                                    'activo': True})
        if not fila.form.is_valid():
            return fila.error('; '.join(_errores_del_form(fila.form)))
        vistos[placa] = fila.numero
        return fila


def _tipos_documento():
    tipos = {}
    for tipo in TipoDocumentoIdentidad.objects.filter(activo=True):
        tipos[clave(tipo.codigo)] = tipo
        tipos[clave(tipo.nombre)] = tipo
    return tipos


def _documento(fila, contexto):
    """(tipo, número) de la fila, reponiendo los ceros que Excel quita. Marca error si el tipo no existe."""
    v = fila.valores
    tipo = _buscar(contexto['tipos'], v['tipo_documento'])
    if tipo is None:
        fila.error(f"Tipo de documento «{v['tipo_documento']}» no existe. Use: "
                   + ', '.join(sorted({t.codigo for t in contexto['tipos'].values()})) + '.')
        return None, None
    numero = v['numero_documento']
    if tipo.solo_digitos and numero.isdigit() and (
            tipo.longitud_min - CEROS_REPONIBLES <= len(numero) < tipo.longitud_min):
        numero = numero.zfill(tipo.longitud_min)  # Excel quitó los ceros iniciales
        v['numero_documento'] = numero
    return tipo, numero


# --- Licencia de conducir (columnas comunes a la carga de conductores y a la de licencias) ---

CAMPOS_LICENCIA_EXCEL = ['licencia_numero', 'licencia_categoria', 'licencia_vencimiento',
                         'licencia_especial_categoria', 'licencia_especial_vencimiento']
LISTAS_LICENCIA = {'categorias': [valor for valor, _ in CATEGORIAS_LICENCIA],
                   'categorias_especiales': [valor for valor, _ in CATEGORIAS_LICENCIA_ESPECIAL]}
ETIQUETAS_CAMBIO = {
    'licencia_numero': 'número de licencia', 'licencia_categoria': 'categoría',
    'licencia_vencimiento': 'vencimiento', 'licencia_especial_categoria': 'licencia especial',
    'licencia_especial_vencimiento': 'vencimiento especial', 'notificar_vencimiento': 'notificaciones',
}


def _columnas_licencia(obligatoria):
    opcional = '' if obligatoria else ' Obligatorio si escribe la licencia.'
    return [
        Columna('licencia_numero', 'licencia_numero', obligatoria,
                "Número de licencia de conducir (ej. Q43028991)." + ('' if obligatoria else ' Opcional.'),
                'Q08123456', tipo='documento'),
        Columna('licencia_categoria', 'licencia_categoria', obligatoria,
                "Categoría de la licencia de camiones: una de la lista." + opcional, 'A-IIIb', lista='categorias'),
        Columna('licencia_vencimiento', 'licencia_vencimiento', obligatoria,
                "Vencimiento de la licencia: DD/MM/AAAA." + opcional, None, tipo='fecha'),
        Columna('licencia_especial_categoria', 'licencia_especial_categoria', False,
                "A-IV si tiene licencia especial para materiales y residuos peligrosos (MAPTEL). Opcional.",
                'A-IV', lista='categorias_especiales'),
        Columna('licencia_especial_vencimiento', 'licencia_especial_vencimiento', False,
                "Vencimiento de la licencia especial: DD/MM/AAAA. Obligatorio si indica A-IV.", None, tipo='fecha'),
    ]


def _categorias(valores):
    """Normaliza las categorías escritas a mano ('A IIIb' -> 'A-IIIb') en los valores de la fila."""
    if valores.get('licencia_categoria'):
        valores['licencia_categoria'] = licencias.normalizar_categoria(valores['licencia_categoria'],
                                                                       CATEGORIAS_LICENCIA)
    if valores.get('licencia_especial_categoria'):
        valores['licencia_especial_categoria'] = licencias.normalizar_categoria(
            valores['licencia_especial_categoria'], CATEGORIAS_LICENCIA_ESPECIAL)


def _ejemplo_con_fechas(tipo, fila):
    """Vencimientos de ejemplo en el futuro (la plantilla se genera al momento)."""
    nombres = [c.nombre for c in tipo.columnas]
    inicio_de_mes = tiempo.hoy().replace(day=1)
    fila[nombres.index('licencia_vencimiento')] = inicio_de_mes + timedelta(days=3 * 365)
    fila[nombres.index('licencia_especial_vencimiento')] = inicio_de_mes + timedelta(days=365)
    return fila


class CargaConductores(TipoCarga):
    clave = 'conductores'
    titulo = 'Conductores'
    singular = 'conductor'
    nota = ("Los documentos que ya existen en el maestro se omiten (no se modifican); para cambiar la licencia "
            "de conductores ya registrados use «Actualizar licencias». Si Excel quitó el cero inicial de un DNI, "
            "se completa solo. La licencia es opcional: si la escribe, indique número, categoría y vencimiento. "
            "Las notificaciones de vencimiento quedan activas.")
    columnas = [
        Columna('tipo_documento', 'tipo_documento', True, "DNI o CE.", 'DNI', lista='tipos_documento'),
        Columna('numero_documento', 'numero_documento', True,
                "DNI: 8 dígitos. CE: 9 a 12 letras o números. Escríbalo como texto para conservar los ceros.",
                '08123456', tipo='documento'),
        Columna('apellido_paterno', 'apellido_paterno', True, "Apellido paterno.", 'PÉREZ'),
        Columna('apellido_materno', 'apellido_materno', False, "Apellido materno (opcional).", 'GÓMEZ'),
        Columna('nombres', 'nombres', True, "Nombres.", 'JUAN CARLOS'),
    ] + _columnas_licencia(obligatoria=False)

    def listas(self):
        return {'tipos_documento': list(TipoDocumentoIdentidad.objects.filter(activo=True)
                                        .values_list('codigo', flat=True)), **LISTAS_LICENCIA}

    def ejemplo(self):
        return _ejemplo_con_fechas(self, super().ejemplo())

    def preparar(self):
        return {'tipos': _tipos_documento(),
                'existentes': set(Conductor.objects.values_list('tipo_documento_id', 'numero_documento'))}

    def validar(self, fila, contexto, vistos, ahora):
        v = fila.valores
        tipo, numero = _documento(fila, contexto)
        if tipo is None:
            return fila
        llave = (tipo.pk, numero)
        if llave in contexto['existentes']:
            return fila.omitir(f"El {tipo.codigo} {numero} ya está registrado; se omite.")
        if llave in vistos:
            return fila.omitir(f"Documento repetido en el archivo (fila {vistos[llave]}); se omite.")
        _categorias(v)
        fila.form = ConductorForm({'tipo_documento': tipo.pk, 'numero_documento': numero,
                                   'apellido_paterno': v['apellido_paterno'],
                                   'apellido_materno': v['apellido_materno'], 'nombres': v['nombres'],
                                   'activo': True, 'notificar_vencimiento': True,
                                   **{c: v.get(c, '') for c in CAMPOS_LICENCIA_EXCEL}})
        if not fila.form.is_valid():
            return fila.error('; '.join(_errores_del_form(fila.form)))
        vistos[llave] = fila.numero
        return fila


class CargaLicencias(TipoCarga):
    """Actualiza la licencia de conductores YA registrados: es la única carga que modifica registros, y
    solo toca la licencia y las notificaciones (los datos personales no se cambian)."""
    clave = 'licencias'
    titulo = 'Licencias'
    singular = 'licencia'
    participio = 'Actualizada'
    verbo_plural = 'actualizaron'
    volver_a = 'Conductores'
    titulo_propio = "Actualizar licencias de conductores"
    nota = ("Actualiza la licencia de conductores que ya existen en el maestro (se buscan por tipo y número de "
            "documento). Número, categoría y vencimiento son obligatorios. Si deja vacías las dos columnas de la "
            "licencia especial, se conserva la que tenga registrada. La columna notificar (SI/NO) es opcional: "
            "vacía conserva el valor actual, y al renovar una licencia las notificaciones se vuelven a activar. "
            "Las filas sin cambios se omiten.")
    columnas = [
        Columna('tipo_documento', 'tipo_documento', True, "DNI o CE del conductor.", 'DNI', lista='tipos_documento'),
        Columna('numero_documento', 'numero_documento', True,
                "Documento de un conductor registrado. Escríbalo como texto para conservar los ceros.",
                '08123456', tipo='documento'),
    ] + _columnas_licencia(obligatoria=True) + [
        Columna('notificar', 'notificar', False,
                "SI o NO: avisar en la campanita antes del vencimiento. Vacío conserva el valor actual.",
                'SI', lista='si_no'),
    ]

    def listas(self):
        return {'tipos_documento': list(TipoDocumentoIdentidad.objects.filter(activo=True)
                                        .values_list('codigo', flat=True)), **LISTAS_LICENCIA, 'si_no': ['SI', 'NO']}

    def ejemplo(self):
        return _ejemplo_con_fechas(self, super().ejemplo())

    def preparar(self):
        return {'tipos': _tipos_documento(),
                'conductores': {(c.tipo_documento_id, c.numero_documento): c
                                for c in Conductor.objects.select_related('tipo_documento')}}

    def validar(self, fila, contexto, vistos, ahora):
        v = fila.valores
        tipo, numero = _documento(fila, contexto)
        if tipo is None:
            return fila
        llave = (tipo.pk, numero)
        conductor = contexto['conductores'].get(llave)
        if conductor is None:
            return fila.error(f"El {tipo.codigo} {numero} no está registrado en Conductores. Regístrelo primero "
                              "(pantalla o carga masiva de conductores).")
        if llave in vistos:
            return fila.error(f"El conductor se repite en el archivo (fila {vistos[llave]}); deje una sola fila.")
        notificar = None
        if v.get('notificar'):
            try:
                notificar = a_si_no(v['notificar'])
            except ValueError as error:
                return fila.error(f"Notificar: {error}.")
        _categorias(v)
        datos = {c: v.get(c, '') for c in CAMPOS_LICENCIA_EXCEL}
        if not datos['licencia_especial_categoria'] and not datos['licencia_especial_vencimiento']:
            # Columnas vacías: se conserva la licencia especial registrada.
            datos['licencia_especial_categoria'] = conductor.licencia_especial_categoria
            datos['licencia_especial_vencimiento'] = conductor.licencia_especial_vencimiento or ''
        elif not datos['licencia_especial_categoria']:
            datos['licencia_especial_categoria'] = CATEGORIAS_LICENCIA_ESPECIAL[0][0]  # solo existe A-IV
        datos['notificar_vencimiento'] = conductor.notificar_vencimiento
        fila.foto_antes = auditoria.foto(conductor)
        fila.form = ConductorForm(datos, instance=conductor, solo_licencia=True)
        if not fila.form.is_valid():
            return fila.error('; '.join(_errores_del_form(fila.form)))
        if notificar is not None:
            fila.form.instance.notificar_vencimiento = notificar  # lo que dice el archivo manda
        _, cambios = auditoria.diferencias(fila.foto_antes, auditoria.foto(fila.form.instance))
        if not cambios:
            return fila.omitir(f"{conductor.nombre_completo}: sin cambios; se omite.")
        vistos[llave] = fila.numero
        fila.mensajes.append(f"{conductor.nombre_completo}: cambia "
                             + ', '.join(ETIQUETAS_CAMBIO.get(c, c) for c in cambios) + '.')
        return fila

    def grabar(self, fila, usuario, motivo):
        conductor = fila.form.save(commit=False)
        conductor.modificado_por = usuario
        conductor.save()
        auditoria.registrar_modificacion(usuario, conductor, fila.foto_antes, motivo=motivo)
        return conductor


class CargaProgramaciones(TipoCarga):
    clave = 'programaciones'
    titulo = 'Programaciones'
    singular = 'programación'
    nota = ("La placa y el conductor deben existir y estar activos en los maestros (cárguelos primero). "
            "Se aplican las mismas reglas que en la pantalla: fecha y hora futuras (hora de Lima), sin domingos ni "
            "feriados, RUC válido, la placa no se repite en la misma fecha y hora, INTERNO solo con transportes "
            "propios y disponibles. En INTERNO la hora es la de SALIDA del CD y decide el turno: desde las 20:00 "
            "hasta las 06:15 del día siguiente es turno noche (de lunes a viernes; el sábado es excepcional; no la "
            "noche de un feriado ni la anterior); las demás, turno día. Una salida de madrugada (00:00 a 06:15) se "
            "carga con la fecha en que sale. La primera ruta INTERNO de una placa en cada turno crea su Pre-Ruta "
            "automática (06:30 el día, 20:00 la noche). Los avisos (capacidad, sábado, mismo conductor en los dos "
            "turnos) no impiden grabar.")
    columnas = [
        Columna('cd', 'cd', True, "Código o nombre del CD: uno de la lista.", 'LURIN', lista='cds'),
        Columna('fecha', 'fecha', True, "Fecha de la cita o de la salida: DD/MM/AAAA (también acepta AAAAMMDD). "
                "Una salida de madrugada del turno noche va con la fecha del día en que sale.", None, tipo='fecha'),
        Columna('hora', 'hora', True, "24 horas: HH:MM (ej. 08:30, 22:00). Clientes y proveedores: hora de llegada. "
                "INTERNO: hora de salida del CD; desde las 20:00 hasta las 06:15 es turno noche.", '08:30',
                tipo='hora'),
        Columna('tipo', 'tipo', True, "CLIENTE, PROVEEDOR o INTERNO.", 'CLIENTE', lista='tipos'),
        Columna('placa', 'placa', True, "Placa registrada en el maestro de transportes.", 'F7J-011'),
        Columna('conductor_documento', 'conductor_documento', True,
                "Número de DNI o CE de un conductor registrado en el maestro.", '08123456', tipo='documento'),
        Columna('ruc_dni', 'ruc_dni', False,
                "RUC (11 dígitos) o DNI (8) del cliente o proveedor. En INTERNO puede quedar vacío.",
                '20100047218', tipo='documento'),
        Columna('razon_social', 'razon_social', False,
                "Razón social del cliente o proveedor. En INTERNO puede quedar vacía.", 'CLIENTE EJEMPLO S.A.C.'),
        Columna('motivo', 'motivo', True, "Motivo de la visita.", 'Recojo de pedido'),
        Columna('observacion', 'observacion', False, "Observación de Despacho (opcional).", ''),
    ]

    def listas(self):
        return {'cds': list(CentroDistribucion.objects.filter(activo=True).values_list('codigo', flat=True)),
                'tipos': [valor for valor, _ in Programacion.Tipo.choices]}

    def ejemplo(self):
        fila = super().ejemplo()
        fila[1] = tiempo.hoy() + timedelta(days=1)
        return fila

    def preparar(self):
        cds = {}
        for cd in CentroDistribucion.objects.filter(activo=True):
            cds[clave(cd.codigo)] = cd
            cds[clave(cd.nombre)] = cd
        conductores = {}
        for conductor in Conductor.objects.select_related('tipo_documento'):
            conductores.setdefault(conductor.numero_documento, []).append(conductor)
        return {'cds': cds, 'tipos': {clave(v): v for v, _ in Programacion.Tipo.choices},
                'transportes': {t.placa_normalizada: t for t in Transporte.objects.select_related('tamano')},
                'conductores': conductores, 'config': ConfiguracionAccesos.obtener(),
                'no_disponibles': jornadas.ids_no_disponibles()}

    def validar(self, fila, contexto, vistos, ahora):
        v = fila.valores
        cd = _buscar(contexto['cds'], v['cd'])
        if cd is None:
            fila.error(f"CD «{v['cd']}» no existe. Use: "
                       + ', '.join(sorted({c.codigo for c in contexto['cds'].values()})) + '.')
        tipo = contexto['tipos'].get(clave(v['tipo']))
        if tipo is None:
            fila.error(f"Tipo «{v['tipo']}» no válido. Use CLIENTE, PROVEEDOR o INTERNO.")
        transporte = contexto['transportes'].get(normalizar_placa(v['placa']))
        if transporte is None:
            fila.error(f"La placa {v['placa']} no está registrada en Transportes.")
        elif not transporte.activo:
            fila.error(f"La placa {v['placa']} está inactiva.")
        elif transporte.pk in contexto['no_disponibles']:
            fila.error(f"La placa {v['placa']} no está disponible (avería registrada): no se puede programar hasta "
                       "que se habilite.")
        conductor = self._conductor(fila, contexto)
        if fila.estado == ERROR:
            return fila

        llave = (transporte.placa_normalizada, v['fecha'], v['hora'])
        if llave in vistos:
            return fila.error(f"La placa {transporte.placa} se repite en el archivo el mismo día y hora "
                              f"(fila {vistos[llave]}).")
        fila.form = ProgramacionForm({
            'cd': cd.pk, 'fecha': v['fecha'], 'hora': v['hora'], 'tipo': tipo, 'transporte': transporte.pk,
            'conductor': conductor.pk, 'documento': v['ruc_dni'], 'razon_social': v['razon_social'],
            'motivo': v['motivo'], 'observacion_despacho': v['observacion'],
        }, ahora=ahora)
        if not fila.form.is_valid():
            return fila.error('; '.join(_errores_del_form(fila.form)))
        vistos[llave] = fila.numero
        avisos = servicios.avisos(fila.form.instance, contexto['config'])
        if avisos:
            fila.estado = AVISO
            fila.mensajes.extend(avisos)
        if tipo == Programacion.Tipo.INTERNO:
            # Informativo: la primera ruta INTERNO de la placa en el turno (día o noche) crea su Pre-Ruta automática.
            config = contexto['config']
            fecha, turno = jornadas.turno_de(fila.form.cleaned_data['fecha'], fila.form.cleaned_data['hora'], config)
            llave_turno = (transporte.pk, fecha, turno)
            existe = llave_turno in contexto.setdefault('jornadas_vistas', set()) or jornadas.jornada_vigente(*llave_turno)
            contexto['jornadas_vistas'].add(llave_turno)
            if turno == jornadas.NOCHE:
                fila.mensajes.append("Pre-Ruta del turno noche: ya existe." if existe else
                                     f"Se creará el turno noche del {fecha:%d/%m} (Pre-Ruta "
                                     f"{config.hora_pre_ruta_noche:%H:%M}).")
            else:
                fila.mensajes.append("Pre-Ruta del día: ya existe." if existe else
                                     f"Se creará la Pre-Ruta automática ({config.hora_pre_ruta:%H:%M}).")
        return fila

    def _conductor(self, fila, contexto):
        numero = fila.valores['conductor_documento']
        candidatos = contexto['conductores'].get(numero, [])
        if not candidatos and numero.isdigit() and 8 - CEROS_REPONIBLES <= len(numero) < 8:
            numero = numero.zfill(8)  # DNI al que Excel le quitó los ceros iniciales
            candidatos = contexto['conductores'].get(numero, [])
            fila.valores['conductor_documento'] = numero
        if not candidatos:
            fila.error(f"El conductor con documento {numero} no está registrado en Conductores.")
            return None
        if len(candidatos) > 1:
            fila.error(f"Hay más de un conductor con el número {numero}; regístrelo desde la pantalla.")
            return None
        if not candidatos[0].activo:
            fila.error(f"El conductor {candidatos[0].nombre_completo} está inactivo.")
            return None
        return candidatos[0]

    def despues_de_grabar(self, fila, objeto):
        # Avisos recalculados ya con las filas anteriores del archivo grabadas (capacidad, misma placa).
        avisos = servicios.avisos(objeto, ConfiguracionAccesos.obtener())
        pre_ruta = [m.replace('Se creará', 'Se creó') for m in fila.mensajes
                    if m.startswith(('Pre-Ruta del día', 'Pre-Ruta del turno noche', 'Se creará la Pre-Ruta',
                                     'Se creará el turno noche'))]
        fila.mensajes = [f"{objeto.codigo} registrada."] + avisos + pre_ruta
        fila.estado = AVISO if avisos else OK


TIPOS = {t.clave: t for t in (CargaTransportes(), CargaConductores(), CargaProgramaciones(), CargaLicencias())}


# --- PLANTILLA ---

def libro_plantilla(tipo):
    """Plantilla .xlsx: hoja «Datos» vacía con encabezados, formatos y listas desplegables,
    hoja «Instrucciones» y hoja «Ejemplo». Devuelve los bytes del archivo."""
    libro = Workbook()
    datos = libro.active
    datos.title = 'Datos'
    relleno = PatternFill(start_color=COLOR_CABECERA, end_color=COLOR_CABECERA, fill_type='solid')
    blanco = Font(color='FFFFFF', bold=True)

    def cabecera(hoja, titulos):
        for indice, titulo in enumerate(titulos, start=1):
            celda = hoja.cell(row=1, column=indice, value=titulo)
            celda.fill, celda.font = relleno, blanco
            celda.alignment = Alignment(horizontal='center')
        hoja.freeze_panes = 'A2'

    cabecera(datos, [c.nombre for c in tipo.columnas])

    listas = tipo.listas()
    if listas:
        hoja_listas = libro.create_sheet('Listas')
        rangos = {}
        for indice, (nombre, valores) in enumerate(listas.items(), start=1):
            letra = get_column_letter(indice)
            hoja_listas.cell(row=1, column=indice, value=nombre)
            for fila, valor in enumerate(valores, start=2):
                hoja_listas.cell(row=fila, column=indice, value=valor)
            rangos[nombre] = f"Listas!${letra}$2:${letra}${max(len(valores), 1) + 1}"
        hoja_listas.sheet_state = 'hidden'

    formatos = {'fecha': 'DD/MM/YYYY', 'hora': 'HH:MM', 'documento': '@', 'texto': '@'}
    for indice, col in enumerate(tipo.columnas, start=1):
        letra = get_column_letter(indice)
        datos.column_dimensions[letra].width = max(len(col.nombre) + 6, 16)
        if col.lista and col.lista in rangos:
            validacion = DataValidation(type='list', formula1=rangos[col.lista], allow_blank=not col.obligatoria,
                                        showErrorMessage=True, errorTitle='Valor no válido',
                                        error='Elija un valor de la lista.')
            validacion.add(f"{letra}2:{letra}{FILAS_CON_FORMATO}")
            datos.add_data_validation(validacion)
        elif col.tipo in ('documento', 'texto', 'fecha', 'hora'):
            for fila in range(2, FILAS_CON_FORMATO + 1):
                datos.cell(row=fila, column=indice).number_format = formatos[col.tipo]

    instrucciones = libro.create_sheet('Instrucciones')
    cabecera(instrucciones, ['Columna', 'Obligatoria', 'Formato y valores'])
    for fila, col in enumerate(tipo.columnas, start=2):
        valores = f" Valores: {', '.join(map(str, listas[col.lista]))}." if col.lista in listas else ''
        instrucciones.append([col.nombre, 'Sí' if col.obligatoria else 'No', col.formato + valores])
    instrucciones.append([])
    instrucciones.append(['Reglas', '', tipo.nota])
    instrucciones.append(['', '', f"Escriba los datos en la hoja «Datos» desde la fila 2, sin cambiar los "
                                  f"encabezados. Máximo {MAX_FILAS} filas por archivo."])
    instrucciones.append(['', '', "Al subir el archivo verá una vista previa fila por fila; nada se graba hasta "
                                  "que confirme. Solo se graban las filas válidas y se informan los errores."])
    instrucciones.column_dimensions['A'].width = 24
    instrucciones.column_dimensions['B'].width = 12
    instrucciones.column_dimensions['C'].width = 110

    ejemplo = libro.create_sheet('Ejemplo')
    cabecera(ejemplo, [c.nombre for c in tipo.columnas])
    ejemplo.append(tipo.ejemplo())
    for indice, col in enumerate(tipo.columnas, start=1):
        letra = get_column_letter(indice)
        ejemplo.column_dimensions[letra].width = max(len(col.nombre) + 6, 16)
        if col.tipo == 'fecha':
            ejemplo.cell(row=2, column=indice).number_format = 'DD/MM/YYYY'
        elif col.tipo == 'documento':
            ejemplo.cell(row=2, column=indice).number_format = '@'

    libro.active = 0
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
