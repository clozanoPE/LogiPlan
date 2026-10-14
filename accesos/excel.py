"""
Exportación a Excel del módulo de Accesos, en memoria (BytesIO), con cabecera naranja (el color de Control de
Accesos; Planificación conserva su verde). Es un helper propio para no acoplar este módulo a `planificacion.views`.
"""
from datetime import date, datetime, time
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_CABECERA = 'EA580C'  # naranja del módulo (igual que las plantillas de carga masiva)
MAX_CARACTERES_CELDA = 32767  # límite de Excel
ANCHO_MAXIMO = 50


def si_no(valor):
    """Booleano como SI / NO (el mismo texto que aceptan las cargas masivas)."""
    return 'SI' if valor else 'NO'


def _valor_seguro(valor):
    if isinstance(valor, str) and len(valor) > MAX_CARACTERES_CELDA:
        return valor[:MAX_CARACTERES_CELDA - 1] + '…'
    return valor


def _largo(valor):
    if valor is None:
        return 0
    if isinstance(valor, (datetime, date)):
        return 12
    if isinstance(valor, time):
        return 6
    return max((len(linea) for linea in str(valor).splitlines()), default=0)


def _indices(encabezados, formatos):
    """`formatos` admite el índice de la columna (0-based) o su encabezado."""
    resultado = {}
    for clave, formato in (formatos or {}).items():
        indice = encabezados.index(clave) if isinstance(clave, str) else clave
        resultado[indice] = formato
    return resultado


def _cabecera(ws, encabezados):
    relleno = PatternFill(start_color=COLOR_CABECERA, end_color=COLOR_CABECERA, fill_type='solid')
    for columna, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=columna, value=encabezado)
        celda.fill = relleno
        celda.font = Font(color='FFFFFF', bold=True)
        celda.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'


def libro_excel(hoja, encabezados, filas, formatos=None, notas=None):
    """Bytes de un .xlsx con una hoja de datos y, si hay `notas`, una hoja «Notas» con una línea por nota.
    `formatos`: {índice o encabezado: number_format}. Las celdas no admiten datetimes con zona: pasar fechas y
    horas ya convertidas a hora de Lima, sin tzinfo."""
    formatos = _indices(encabezados, formatos)
    libro = Workbook()
    ws = libro.active
    ws.title = hoja
    _cabecera(ws, encabezados)
    anchos = [len(e) + 4 for e in encabezados]
    for numero_fila, fila in enumerate(filas, start=2):
        for indice, valor in enumerate(fila):
            celda = ws.cell(row=numero_fila, column=indice + 1, value=_valor_seguro(valor))
            if celda.data_type == 'f':
                # Un texto escrito por un usuario que empieza con «=» se guarda como texto, nunca como fórmula.
                celda.data_type = 's'
            if indice in formatos:
                celda.number_format = formatos[indice]
            anchos[indice] = max(anchos[indice], _largo(valor) + 2)
    for columna, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(columna)].width = min(max(ancho, 12), ANCHO_MAXIMO)

    if notas:
        hoja_notas = libro.create_sheet('Notas')
        _cabecera(hoja_notas, ['Notas'])
        for numero_fila, nota in enumerate(notas, start=2):
            hoja_notas.cell(row=numero_fila, column=1, value=nota).alignment = Alignment(wrap_text=True)
        hoja_notas.column_dimensions['A'].width = 120

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def respuesta_excel(nombre_archivo, hoja, encabezados, filas, formatos=None, notas=None):
    respuesta = HttpResponse(libro_excel(hoja, encabezados, filas, formatos, notas),
                             content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return respuesta
