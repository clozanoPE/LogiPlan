from datetime import datetime, time
from io import BytesIO

from django.core.exceptions import ValidationError
from django.urls import reverse
from openpyxl import load_workbook

from .. import servicios, tiempo
from ..models import RegistroAuditoria, Tramo
from .test_porteria import PorteriaBase
from .test_programacion import HOY


class CorreccionesTests(PorteriaBase):

    def tramo_cerrado(self, programacion=None):
        tramo = self.entrar_en(programacion or self.programar(9, 0))
        self.mover_reloj(HOY, 9, 40)
        return self.salir(tramo)

    def test_corregir_recalcula_y_audita(self):
        tramo = self.tramo_cerrado()
        nueva_entrada = tiempo.combinar(HOY, time(9, 15))
        servicios.corregir_tramo(
            tramo, self.admin_accesos, motivo='Hora mal registrada', ahora=tiempo.ahora(), entrada_en=nueva_entrada,
            salida_en=tramo.salida_en, km_entrada=125500, obs_entrada='Ninguna', obs_salida='Ninguna',
            programacion=tramo.programacion, conductor=tramo.conductor)
        tramo.refresh_from_db()
        self.assertTrue(tramo.corregido)
        self.assertEqual(tramo.minutos_permanencia, 25)
        self.assertEqual(tramo.desviacion_llegada_min, 15)
        self.assertEqual(tramo.hora_local_entrada, time(9, 15))
        registro = RegistroAuditoria.objects.get(entidad='Tramo', objeto_id=tramo.pk, accion='CORREGIR')
        self.assertEqual(registro.motivo, 'Hora mal registrada')
        self.assertEqual(registro.datos_antes['km_entrada'], 125430)

    def test_corregir_exige_motivo_y_horas_coherentes(self):
        tramo = self.tramo_cerrado()
        base = dict(ahora=tiempo.ahora(), salida_en=tramo.salida_en, km_entrada=1, obs_entrada='', obs_salida='',
                    programacion=tramo.programacion, conductor=tramo.conductor)
        with self.assertRaises(ValidationError):
            servicios.corregir_tramo(tramo, self.admin_accesos, motivo='', entrada_en=tramo.entrada_en, **base)
        with self.assertRaises(ValidationError):
            servicios.corregir_tramo(tramo, self.admin_accesos, motivo='x',
                                     entrada_en=tiempo.combinar(HOY, time(9, 50)), **base)

    def test_reasignar_a_otra_programacion_de_la_placa(self):
        tramo = self.tramo_cerrado()
        correcta = self.programar(8, 50)
        self.assertIn(correcta, servicios.programaciones_para_corregir(tramo))
        servicios.corregir_tramo(
            tramo, self.admin_accesos, motivo='Vigilancia eligió otra', ahora=tiempo.ahora(),
            entrada_en=tramo.entrada_en, salida_en=tramo.salida_en, km_entrada=tramo.km_entrada,
            obs_entrada='', obs_salida='', programacion=correcta, conductor=tramo.conductor)
        tramo.refresh_from_db()
        self.assertEqual(tramo.programacion, correcta)
        self.assertEqual(tramo.desviacion_llegada_min, 10)

    def test_anular_deja_la_programacion_pendiente_otra_vez(self):
        programacion = self.programar(9, 0)
        tramo = self.entrar_en(programacion)
        servicios.anular_tramo(tramo, self.admin_accesos, motivo='Entrada equivocada')
        tramo.refresh_from_db()
        self.assertTrue(tramo.esta_anulado)
        self.assertFalse(tramo.esta_abierto)
        programacion.refresh_from_db()
        self.assertEqual(programacion.estado, 'PROGRAMADA')
        self.assertEqual(servicios.estado_visible(programacion, HOY), servicios.PENDIENTE)
        nuevo = self.entrar_en(programacion)  # la misma programación puede volver a usarse
        self.assertTrue(nuevo.vigente)
        self.assertEqual(Tramo.objects.filter(programacion=programacion).count(), 2)  # el anulado no se borra
        self.assertTrue(RegistroAuditoria.objects.filter(entidad='Tramo', objeto_id=tramo.pk, accion='ANULAR').exists())

    def test_cierre_administrativo_de_salida_olvidada(self):
        tramo = self.entrar_en(self.programar(9, 0))
        self.mover_reloj(HOY, 18, 0)
        with self.assertRaises(ValidationError):
            servicios.cierre_administrativo(tramo, self.admin_accesos, salida_en=tiempo.combinar(HOY, time(19, 0)),
                                            motivo='x', ahora=tiempo.ahora())
        servicios.cierre_administrativo(tramo, self.admin_accesos, salida_en=tiempo.combinar(HOY, time(11, 0)),
                                        motivo='Vigilancia olvidó la salida', ahora=tiempo.ahora())
        tramo.refresh_from_db()
        self.assertTrue(tramo.cierre_administrativo)
        self.assertEqual(tramo.minutos_permanencia, 120)
        self.assertIn('Cierre administrativo', tramo.obs_salida)


class VistasDeControlTests(PorteriaBase):

    def test_lectura_para_administrador_despacho_y_consulta(self):
        tramo = self.entrar_en(self.programar(9, 0))
        indicadores = [reverse('accesos:tablero_hoy'), reverse('accesos:dashboard')]
        urls = [reverse('accesos:tramos_lista'), reverse('accesos:tramos_exportar'),
                reverse('accesos:tramo_detalle', args=[tramo.pk])] + indicadores
        for usuario in (self.admin_accesos, self.despacho, self.consulta):
            self.entrar(usuario)
            for url in urls:
                self.assertEqual(self.client.get(url).status_code, 200, (usuario.username, url))
        # Vigilancia también lee «Hoy», el Dashboard y los tramos de su CD (sin acciones de corrección).
        self.entrar(self.vigilante)
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.assertNotContains(self.client.get(urls[2]), 'Corregir datos')

    def test_vigilancia_solo_ve_los_tramos_de_sus_cd(self):
        propio = self.entrar_en(self.programar(9, 0))
        ajeno = self.entrar_en(self.programar(9, 0, transporte=self.externo, cd=self.punta_negra))
        self.entrar(self.vigilante)  # autorizado solo en Lurín
        lista = self.client.get(reverse('accesos:tramos_lista'))
        self.assertEqual([t.pk for t in lista.context['page_obj']], [propio.pk])
        self.assertEqual(list(lista.context['cds']), [self.lurin])
        # Aunque pida el otro CD en la URL, no lo ve.
        filtrada = self.client.get(reverse('accesos:tramos_lista'), {'cd': self.punta_negra.pk})
        self.assertEqual(len(filtrada.context['page_obj']), 0)
        self.assertEqual(self.client.get(reverse('accesos:tramo_detalle', args=[ajeno.pk])).status_code, 404)
        libro = load_workbook(BytesIO(self.client.get(reverse('accesos:tramos_exportar')).content))
        self.assertEqual(libro.active.max_row, 2)  # cabecera + su tramo
        # En «Hoy» la placa del otro CD aparece, pero sin enlace al detalle.
        hoy = self.client.get(reverse('accesos:tablero_hoy'))
        self.assertContains(hoy, reverse('accesos:tramo_detalle', args=[propio.pk]))
        self.assertNotContains(hoy, reverse('accesos:tramo_detalle', args=[ajeno.pk]))
        # El código de la programación no enlaza a una pantalla que Vigilancia no puede abrir.
        detalle = self.client.get(reverse('accesos:tramo_detalle', args=[propio.pk]))
        self.assertNotContains(detalle, reverse('accesos:programacion_detalle', args=[propio.programacion_id]))
        # Los demás roles ven ambos.
        self.entrar(self.consulta)
        self.assertEqual(len(self.client.get(reverse('accesos:tramos_lista')).context['page_obj']), 2)

    def test_correcciones_solo_administrador(self):
        tramo = self.entrar_en(self.programar(9, 0))
        for usuario in (self.despacho, self.consulta, self.vigilante):
            self.entrar(usuario)
            self.assertEqual(self.client.get(reverse('accesos:tramo_corregir', args=[tramo.pk])).status_code, 403)
            self.assertEqual(self.client.post(reverse('accesos:tramo_anular', args=[tramo.pk]),
                                              {'anular-motivo': 'x'}).status_code, 403)
        self.entrar(self.admin_accesos)
        self.assertEqual(self.client.get(reverse('accesos:auditoria_lista')).status_code, 200)
        self.assertContains(self.client.get(reverse('accesos:tramo_detalle', args=[tramo.pk])), 'Cierre administrativo')

    def test_correccion_desde_la_pantalla_en_hora_de_lima(self):
        tramo = self.entrar_en(self.programar(9, 0))
        self.mover_reloj(HOY, 9, 40)
        tramo = self.salir(tramo)
        self.entrar(self.admin_accesos)
        formulario = self.client.get(reverse('accesos:tramo_corregir', args=[tramo.pk]))
        self.assertContains(formulario, 'value="2026-10-07T09:00"')  # se muestra en hora de Lima, no en UTC
        self.client.post(reverse('accesos:tramo_corregir', args=[tramo.pk]), {
            'programacion': tramo.programacion_id, 'entrada_en': '2026-10-07T09:05', 'salida_en': '2026-10-07T09:40',
            'km_entrada': 125430, 'conductor': tramo.conductor_id, 'obs_entrada': '', 'obs_salida': '',
            'motivo': 'Ajuste'})
        tramo.refresh_from_db()
        self.assertEqual(tiempo.a_lima(tramo.entrada_en).time(), time(9, 5))
        self.assertEqual(tramo.minutos_permanencia, 35)

    def test_cierre_y_anulacion_desde_la_pantalla(self):
        tramo = self.entrar_en(self.programar(9, 0))
        self.mover_reloj(HOY, 12, 0)
        self.entrar(self.admin_accesos)
        self.client.post(reverse('accesos:tramo_cierre', args=[tramo.pk]),
                         {'cierre-salida_en': '2026-10-07T10:30', 'cierre-motivo': 'Olvido'})
        tramo.refresh_from_db()
        self.assertEqual(tramo.minutos_permanencia, 90)
        self.client.post(reverse('accesos:tramo_anular', args=[tramo.pk]), {'anular-motivo': 'Duplicado'})
        tramo.refresh_from_db()
        self.assertTrue(tramo.esta_anulado)

    def test_exportar_tramos_a_excel(self):
        tramo = self.entrar_en(self.programar(9, 0))
        self.mover_reloj(HOY, 9, 30)
        self.salir(tramo, obs='Sin novedad')
        self.entrar(self.consulta)
        respuesta = self.client.get(reverse('accesos:tramos_exportar'), {'desde': '2026-10-07', 'hasta': '2026-10-07'})
        hoja = load_workbook(BytesIO(respuesta.content)).active
        self.assertEqual(hoja.max_row, 2)
        fila = [celda.value for celda in hoja[2]]
        self.assertEqual(fila[6], 'F7J-011')
        self.assertEqual(fila[4], 30)
        self.assertEqual(fila[19], 'Sin novedad')

    def test_tablero_hoy_cuenta_dentro_y_atendidas(self):
        self.entrar_en(self.programar(9, 0))
        self.programar(11, 0)
        self.entrar(self.consulta)
        centros = {c['cd'].codigo: c for c in self.client.get(reverse('accesos:tablero_hoy')).context['centros']}
        self.assertEqual(centros['LURIN']['dentro'], 1)
        self.assertEqual(centros['LURIN']['atendidas'], 1)
        self.assertEqual(centros['LURIN']['pendientes'], 1)
