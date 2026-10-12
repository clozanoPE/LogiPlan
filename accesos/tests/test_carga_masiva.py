from datetime import date, datetime, time
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .. import carga_masiva
from ..models import Conductor, Programacion, RegistroAuditoria, Transporte
from .test_programacion import HOY, MANANA, RUC_DARYZA, ProgramacionBase


def excel(encabezados, filas, nombre='carga.xlsx'):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Datos'
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return SimpleUploadedFile(nombre, buffer.getvalue())


COLUMNAS_PROG = ['cd', 'fecha', 'hora', 'tipo', 'placa', 'conductor_documento', 'ruc_dni', 'razon_social',
                 'motivo', 'observacion']


class ConversionDeCeldasTests(SimpleTestCase):

    def test_fechas_en_los_formatos_que_produce_excel(self):
        esperado = date(2026, 10, 8)
        for valor in (datetime(2026, 10, 8), esperado, '08/10/2026', '2026-10-08', '20261008', 20261008,
                      20261008.0, 46303):
            self.assertEqual(carga_masiva.a_fecha(valor), esperado, valor)
        with self.assertRaises(ValueError):
            carga_masiva.a_fecha('31/02/2026')

    def test_horas_en_los_formatos_que_produce_excel(self):
        for valor in (time(8, 30), datetime(1900, 1, 1, 8, 30), 0.3541666, '08:30', '8:30', '08:30:00', '8.30'):
            self.assertEqual(carga_masiva.a_hora(valor), time(8, 30), valor)
        with self.assertRaises(ValueError):
            carga_masiva.a_hora('25:00')

    def test_si_no_y_documentos(self):
        self.assertTrue(carga_masiva.a_si_no('Sí'))
        self.assertFalse(carga_masiva.a_si_no(''))
        with self.assertRaises(ValueError):
            carga_masiva.a_si_no('tal vez')
        self.assertEqual(carga_masiva.a_documento(20100047218.0), '20100047218')
        self.assertEqual(carga_masiva.clave(' Razón Social '), 'razon_social')


class CargaBase(ProgramacionBase):

    def setUp(self):
        super().setUp()
        self.entrar(self.despacho)

    def url(self, tipo):
        return reverse('accesos:carga_masiva', args=[tipo])

    def previsualizar(self, tipo, archivo):
        return self.client.post(self.url(tipo), {'accion': 'previsualizar', 'archivo': archivo})

    def confirmar(self, tipo):
        return self.client.post(self.url(tipo), {'accion': 'confirmar'})

    def estados(self, respuesta):
        return [(f.numero, f.estado) for f in respuesta.context['filas']]


class PlantillasTests(CargaBase):

    def test_plantillas_con_hojas_listas_y_desplegables(self):
        for clave, tipo in carga_masiva.TIPOS.items():
            respuesta = self.client.get(reverse('accesos:carga_plantilla', args=[clave]))
            self.assertEqual(respuesta.status_code, 200)
            libro = load_workbook(BytesIO(respuesta.content))
            self.assertEqual(libro.sheetnames[0], 'Datos')
            self.assertIn('Instrucciones', libro.sheetnames)
            self.assertIn('Ejemplo', libro.sheetnames)
            encabezados = [c.value for c in libro['Datos'][1]]
            self.assertEqual(encabezados, [c.nombre for c in tipo.columnas])
            self.assertEqual(libro['Datos'].max_row - 1, carga_masiva.MAX_FILAS)  # filas con formato, sin datos
            self.assertEqual(libro['Listas'].sheet_state, 'hidden')
            self.assertTrue(libro['Datos'].data_validations.dataValidation)
        # La plantilla vacía se puede subir tal cual: el sistema avisa que no hay filas con datos.
        plantilla = self.client.get(reverse('accesos:carga_plantilla', args=['transportes'])).content
        respuesta = self.previsualizar('transportes', SimpleUploadedFile('Plantilla.xlsx', plantilla))
        self.assertRedirects(respuesta, self.url('transportes'))

    def test_solo_administrador_y_despacho(self):
        for usuario in (self.consulta, self.vigilante, self.comercial):
            self.entrar(usuario)
            self.assertEqual(self.client.get(self.url('transportes')).status_code, 403)
            self.assertEqual(self.client.get(reverse('accesos:carga_plantilla', args=['transportes'])).status_code,
                             403)
        for usuario in (self.admin_accesos, self.planeamiento):  # Planeamiento = Despacho
            self.entrar(usuario)
            self.assertEqual(self.client.get(self.url('programaciones')).status_code, 200)
        self.assertEqual(self.client.get(self.url('otra-cosa')).status_code, 404)

    def test_archivos_invalidos(self):
        mensajes = lambda r: [str(m) for m in r.context['messages']]  # noqa: E731
        respuesta = self.client.post(self.url('transportes'), {
            'accion': 'previsualizar', 'archivo': SimpleUploadedFile('datos.csv', b'placa\nABC123')}, follow=True)
        self.assertIn('.xlsx', mensajes(respuesta)[0])
        respuesta = self.client.post(self.url('transportes'), {
            'accion': 'previsualizar', 'archivo': excel(['placa'], [['ABC-999']])}, follow=True)
        self.assertIn('Faltan columnas', mensajes(respuesta)[0])
        muchas = excel(['placa', 'tamano'], [[f'AB{i:04d}', 'x'] for i in range(carga_masiva.MAX_FILAS + 1)])
        respuesta = self.client.post(self.url('transportes'), {'accion': 'previsualizar', 'archivo': muchas},
                                     follow=True)
        self.assertIn('más de', mensajes(respuesta)[0])
        respuesta = self.client.post(self.url('transportes'), {'accion': 'confirmar'}, follow=True)
        self.assertIn('ya no está disponible', mensajes(respuesta)[0])


class CargaTransportesTests(CargaBase):

    def test_vista_previa_no_graba_y_confirmar_graba_solo_las_validas(self):
        archivo = excel(['Placa', 'Tamaño', 'Interno'], [
            ['bcd-456', 'camion grande', 'SI'],      # válida (tamaño sin tilde ni mayúsculas)
            ['F7J 011', 'Camión grande', ''],          # ya existe: se omite
            ['BCD456', 'Camión grande', 'NO'],         # repetida en el archivo: se omite
            ['XY1', 'Camión grande', ''],              # placa corta: error
            ['QWE-789', 'Gigante', ''],                # tamaño inexistente: error
            [None, None, None],                        # fila vacía: se ignora
            ['RTY-321', 'Camión grande', 'quizá'],     # interno inválido: error
        ], nombre='flota.xlsx')
        respuesta = self.previsualizar('transportes', archivo)
        self.assertEqual(self.estados(respuesta), [(2, 'OK'), (3, 'OMITIDA'), (4, 'OMITIDA'), (5, 'ERROR'),
                                                   (6, 'ERROR'), (8, 'ERROR')])
        self.assertContains(respuesta, 'Grabar 1 fila válida')
        self.assertFalse(Transporte.objects.filter(placa_normalizada='BCD456').exists())

        respuesta = self.confirmar('transportes')
        nuevo = Transporte.objects.get(placa_normalizada='BCD456')
        self.assertTrue(nuevo.es_interno)
        self.assertEqual(nuevo.creado_por, self.despacho)
        self.assertEqual(respuesta.context['conteo']['grabadas'], 1)
        registro = RegistroAuditoria.objects.get(entidad='Transporte', objeto_id=nuevo.pk)
        self.assertEqual(registro.motivo, 'Carga masiva: flota.xlsx')
        # La vista previa se consume: no se puede confirmar dos veces.
        self.confirmar('transportes')
        self.assertEqual(Transporte.objects.filter(placa_normalizada='BCD456').count(), 1)

    def test_sin_tamanos_registrados_lo_explica(self):
        from ..models import TamanoVehiculo
        TamanoVehiculo.objects.update(activo=False)
        respuesta = self.previsualizar('transportes', excel(['placa', 'tamano'], [['ZZZ-222', 'Grande']]))
        self.assertIn('No hay tamaños de vehículo registrados', respuesta.context['filas'][0].mensajes[0])

    def test_confirmar_revalida_contra_lo_registrado_despues_de_la_vista_previa(self):
        self.previsualizar('transportes', excel(['placa', 'tamano'], [['ZZZ-111', 'Camión grande']]))
        Transporte.objects.create(placa='ZZZ111', tamano=self.tamano, creado_por=self.admin_accesos)
        respuesta = self.confirmar('transportes')
        self.assertEqual(self.estados(respuesta), [(2, 'OMITIDA')])
        self.assertEqual(Transporte.objects.filter(placa_normalizada='ZZZ111').count(), 1)


class CargaConductoresTests(CargaBase):

    def test_dni_sin_ceros_existentes_y_errores(self):
        archivo = excel(['tipo_documento', 'numero_documento', 'apellido_paterno', 'apellido_materno', 'nombres'], [
            ['DNI', 8123456, 'Ramos', '', 'Luis'],          # Excel quitó el 0: se completa a 08123456
            ['dni', '12345678', 'Otro', '', 'Igual'],       # ya existe: se omite
            ['CE', 'AB12', 'Silva', 'Paz', 'Ana'],          # CE corto: error
            ['PASAPORTE', '123456789', 'X', '', 'Y'],       # tipo inexistente: error
            ['DNI', '87654321', '', '', 'Sin apellido'],    # falta apellido: error
            ['DNI', 123, 'Corto', '', 'Ana'],               # demasiado corto: no se rellena, es error
        ])
        respuesta = self.previsualizar('conductores', archivo)
        self.assertEqual(self.estados(respuesta), [(2, 'OK'), (3, 'OMITIDA'), (4, 'ERROR'), (5, 'ERROR'),
                                                   (6, 'ERROR'), (7, 'ERROR')])
        self.assertIn('8 caracteres', ' '.join(respuesta.context['filas'][-1].mensajes))
        self.confirmar('conductores')
        nuevo = Conductor.objects.get(numero_documento='08123456')
        self.assertEqual(nuevo.apellido_paterno, 'RAMOS')
        self.assertEqual(Conductor.objects.count(), 2)


class CargaProgramacionesTests(CargaBase):

    def setUp(self):
        super().setUp()
        self.inactivo = Conductor.objects.create(tipo_documento=self.dni, numero_documento='11112222',
                                                 apellido_paterno='BAJA', nombres='JOSE', activo=False,
                                                 creado_por=self.admin_accesos)

    def fila(self, **cambios):
        datos = {'cd': 'LURIN', 'fecha': datetime(2026, 10, 8), 'hora': time(8, 30), 'tipo': 'CLIENTE',
                 'placa': 'ABC-123', 'conductor_documento': 12345678, 'ruc_dni': 20100047218,
                 'razon_social': 'Cliente Uno', 'motivo': 'Recojo', 'observacion': ''}
        datos.update(cambios)
        return [datos[c] for c in COLUMNAS_PROG]

    def test_valida_con_las_reglas_del_formulario_y_graba_solo_las_validas(self):
        programada = self.programar(10, 0, transporte=self.externo, fecha=MANANA)
        archivo = excel(COLUMNAS_PROG, [
            self.fila(),                                                     # 2 válida
            self.fila(cd='Punta Negra', fecha='08/10/2026', hora='09:15', tipo='interno', placa='f7j011',
                      ruc_dni='', razon_social=''),                         # 3 INTERNO: RUC de Daryza
            self.fila(hora='08:30'),                                         # 4 repetida en el archivo
            self.fila(fecha=HOY, hora='08:00'),                              # 5 hora ya pasada (son las 9:00)
            self.fila(ruc_dni='20100047219', hora='11:00'),                  # 6 RUC inválido
            self.fila(placa='NOE-000', hora='12:00'),                        # 7 placa no registrada
            self.fila(conductor_documento='11112222', hora='13:00'),         # 8 conductor inactivo
            self.fila(hora='10:00'),                                         # 9 ya programada en la base
            self.fila(tipo='INTERNO', hora='14:00', ruc_dni='', razon_social=''),  # 10 aviso R7 (no interno)
            self.fila(cd='CALLAO', tipo='OTRO', hora='15:00'),               # 11 CD y tipo inválidos
            self.fila(fecha=20261008, hora=0.6875),                          # 12 AAAAMMDD y hora como fracción (16:30)
            self.fila(hora='25:99'),                                         # 13 hora inválida
        ], nombre='semana.xlsx')
        respuesta = self.previsualizar('programaciones', archivo)
        estados = dict(self.estados(respuesta))
        self.assertEqual(estados, {2: 'OK', 3: 'OK', 4: 'ERROR', 5: 'ERROR', 6: 'ERROR', 7: 'ERROR', 8: 'ERROR',
                                   9: 'ERROR', 10: 'AVISO', 11: 'ERROR', 12: 'OK', 13: 'ERROR'})
        mensajes = {f.numero: ' '.join(f.mensajes) for f in respuesta.context['filas']}
        self.assertIn('fila 2', mensajes[4])
        self.assertIn(programada.codigo, mensajes[9])
        self.assertIn('no está registrada', mensajes[7])
        self.assertIn('inactivo', mensajes[8])
        self.assertIn('CALLAO', mensajes[11])
        self.assertIn('OTRO', mensajes[11])
        self.assertEqual(Programacion.objects.count(), 1)  # la vista previa no graba

        respuesta = self.confirmar('programaciones')
        self.assertEqual(respuesta.context['conteo']['grabadas'], 4)
        self.assertEqual(Programacion.objects.count(), 5)
        interna = Programacion.objects.get(transporte=self.f7j)
        self.assertEqual((interna.cd, interna.documento, interna.razon_social),
                         (self.punta_negra, RUC_DARYZA, 'Daryza S.A.C.'))
        self.assertTrue(Programacion.objects.filter(transporte=self.externo, hora=time(16, 30)).exists())
        nuevas = Programacion.objects.exclude(pk=programada.pk)
        self.assertTrue(all(p.creado_por == self.despacho and p.estado == 'PROGRAMADA' for p in nuevas))
        self.assertEqual(RegistroAuditoria.objects.filter(entidad='Programacion', motivo='Carga masiva: semana.xlsx')
                         .count(), 4)
        # El resultado informa el código de cada programación registrada y los errores.
        self.assertContains(respuesta, nuevas.order_by('pk').first().codigo)
        self.assertContains(respuesta, 'no está registrada')

    def test_avisos_de_capacidad_con_las_filas_del_mismo_archivo(self):
        # Punta Negra admite 3 a la vez: la cuarta del archivo a la misma hora genera aviso al grabar.
        placas = ['AAA-111', 'AAA-222', 'AAA-333', 'AAA-444']
        for placa in placas:
            Transporte.objects.create(placa=placa, tamano=self.tamano, creado_por=self.admin_accesos)
        archivo = excel(COLUMNAS_PROG, [self.fila(cd='PUNTA_NEGRA', placa=p) for p in placas])
        self.previsualizar('programaciones', archivo)
        respuesta = self.confirmar('programaciones')
        self.assertEqual(self.estados(respuesta), [(2, 'OK'), (3, 'OK'), (4, 'OK'), (5, 'AVISO')])
        self.assertEqual(Programacion.objects.count(), 4)
