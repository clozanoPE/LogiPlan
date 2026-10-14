"""
Exportación a Excel de las listas: programaciones (agenda y Portería, día o rango), transportes, conductores
(compatible con «Actualizar licencias»), disponibilidad de flota y auditoría. Mismos filtros y permisos que
cada pantalla, cabecera naranja y textos que empiezan con «=» guardados como texto.
"""
from datetime import date, timedelta
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from .. import carga_masiva, consultas, jornadas, servicios, tiempo
from ..models import Conductor, RegistroAuditoria, TipoIncidencia, Transporte
from .test_carga_masiva import CargaBase
from .test_porteria import PorteriaBase
from .test_programacion import HOY, MANANA
from .utilidades import crear_usuario


def libro(respuesta):
    return load_workbook(BytesIO(respuesta.content))


def filas(hoja):
    """[{encabezado: valor}] de las filas de datos."""
    encabezados = [c.value for c in hoja[1]]
    return [dict(zip(encabezados, [c.value for c in fila])) for fila in hoja.iter_rows(min_row=2)]


class ProgramacionesExportarTests(PorteriaBase):

    def setUp(self):
        super().setUp()
        self.mover_reloj(HOY, 9, 0)

    def test_dia_con_los_filtros_de_la_pantalla_y_lo_real(self):
        atendida = self.programar(9, 0)
        self.programar(11, 0, transporte=self.externo)
        tramo = self.entrar_en(atendida, km=125430)
        self.mover_reloj(HOY, 9, 40)
        self.salir(tramo)
        self.entrar(self.despacho)
        respuesta = self.client.get(reverse('accesos:programacion_exportar'), {'fecha': HOY.isoformat()})
        self.assertEqual(respuesta['Content-Disposition'], 'attachment; filename="Programaciones_20261007.xlsx"')
        hoja = libro(respuesta)['Programaciones']
        datos = filas(hoja)
        self.assertEqual([d['Placa'] for d in datos], ['F7J-011', 'ABC-123'])
        primera = datos[0]
        self.assertEqual((primera['Código'], primera['Estado'], primera['La hora es de']),
                         (atendida.codigo, 'Cerrada', 'Llegada'))
        self.assertEqual((primera['Permanencia (min)'], primera['Desviación (min)'], primera['Puntualidad'],
                          primera['Kilometraje'], primera['Clase del tramo']), (40, 0, 'Puntual', 125430, 'Programado'))
        self.assertEqual(primera['Hora entrada'].strftime('%H:%M'), '09:00')
        self.assertEqual(datos[1]['Estado'], 'Pendiente')
        self.assertIsNone(datos[1]['Fecha entrada'])
        self.assertEqual(hoja['A1'].fill.start_color.rgb[-6:], 'EA580C')
        self.assertIn('Notas', libro(respuesta).sheetnames)
        # Los filtros de la pantalla se aplican igual (placa).
        respuesta = self.client.get(reverse('accesos:programacion_exportar'), {'fecha': HOY.isoformat(), 'q': 'abc'})
        self.assertEqual([d['Placa'] for d in filas(libro(respuesta).active)], ['ABC-123'])

    def test_rango_de_hasta_31_dias(self):
        self.programar(10, 0)
        self.programar(10, 0, fecha=MANANA)
        self.programar(10, 0, fecha=HOY + timedelta(days=40))
        self.entrar(self.consulta)  # Consulta también descarga (solo lectura)
        url = reverse('accesos:programacion_exportar')
        respuesta = self.client.get(url, {'desde': HOY.isoformat(), 'hasta': MANANA.isoformat()})
        self.assertEqual(respuesta['Content-Disposition'],
                         'attachment; filename="Programaciones_20261007_20261008.xlsx"')
        self.assertEqual(len(filas(libro(respuesta).active)), 2)
        respuesta = self.client.get(url, {'desde': HOY.isoformat(), 'hasta': (HOY + timedelta(days=60)).isoformat()})
        datos = filas(libro(respuesta).active)
        self.assertEqual(len(datos), 2)  # el día 40 queda fuera del máximo de 31 días
        notas = [c.value for c in libro(respuesta)['Notas']['A']]
        self.assertIn('superaba 31 días', notas[1])
        self.assertEqual(consultas.rango_exportar({'desde': '2026-10-31', 'hasta': '2026-10-01'}, HOY),
                         (date(2026, 10, 1), date(2026, 10, 31), False))

    def test_permisos_y_cd_de_vigilancia(self):
        self.programar(10, 0)
        self.programar(10, 30, cd=self.punta_negra, transporte=self.externo)
        url = reverse('accesos:programacion_exportar')
        for usuario in (self.vigilante, self.comercial, self.sin_grupo):
            self.entrar(usuario)
            self.assertIn(self.client.get(url).status_code, (302, 403), usuario.username)
        self.entrar(self.vigilante)
        # Aunque pida el CD de Punta Negra, solo descarga el suyo (Lurín).
        respuesta = self.client.get(reverse('accesos:porteria_programaciones_exportar'),
                                    {'cd': self.punta_negra.pk})
        self.assertEqual([d['CD'] for d in filas(libro(respuesta).active)], [str(self.lurin)])
        self.entrar(self.despacho)
        self.assertEqual(self.client.get(reverse('accesos:porteria_programaciones_exportar')).status_code, 403)

    def test_botones_en_las_pantallas(self):
        self.entrar(self.consulta)
        pagina = self.client.get(reverse('accesos:programacion_agenda'))
        self.assertContains(pagina, reverse('accesos:programacion_exportar'))
        self.assertContains(pagina, 'Exportar un rango')
        self.entrar(self.vigilante)
        self.assertContains(self.client.get(reverse('accesos:porteria_programaciones')),
                            reverse('accesos:porteria_programaciones_exportar'))

    def test_texto_con_formula_queda_como_texto(self):
        programacion = self.programar(10, 0)
        programacion.motivo = '=HYPERLINK("http://ejemplo","clic")'
        programacion.save()
        self.entrar(self.despacho)
        hoja = libro(self.client.get(reverse('accesos:programacion_exportar'))).active
        columna = [c.value for c in hoja[1]].index('Motivo') + 1
        celda = hoja.cell(row=2, column=columna)
        self.assertEqual((celda.value, celda.data_type), ('=HYPERLINK("http://ejemplo","clic")', 's'))


class MaestrosExportarTests(CargaBase):

    def test_transportes_todos_los_registros_de_la_busqueda(self):
        for i in range(20):
            Transporte.objects.create(placa=f'BUS-{i:03d}', tamano=self.tamano, creado_por=self.admin_accesos)
        Transporte.objects.filter(pk=self.f7j.pk).update(km_actual=130151)
        jornadas.registrar_no_disponibilidad(transporte_id=self.f7j.pk,
                                             tipo_id=TipoIncidencia.objects.get(nombre='Avería mecánica').pk,
                                             descripcion='Frenos', usuario=self.despacho, ahora=tiempo.ahora())
        respuesta = self.client.get(reverse('accesos:transportes_exportar'))
        datos = filas(libro(respuesta)['Transportes'])
        self.assertEqual(len(datos), 22)  # no solo los 15 de la primera página
        f7j = next(d for d in datos if d['Placa'] == 'F7J-011')
        self.assertEqual((f7j['Interno'], f7j['Activo'], f7j['Disponibilidad'], f7j['Último km']),
                         ('SI', 'SI', 'No disponible', 130151))
        respuesta = self.client.get(reverse('accesos:transportes_exportar'), {'q': 'bus-01'})
        self.assertEqual(len(filas(libro(respuesta).active)), 10)
        self.assertContains(self.client.get(reverse('accesos:transportes_lista'), {'q': 'bus'}),
                            reverse('accesos:transportes_exportar') + '?q=bus')

    def test_solo_administrador_y_despacho(self):
        for nombre in ('transportes_exportar', 'conductores_exportar'):
            self.entrar(self.consulta)
            self.assertEqual(self.client.get(reverse(f'accesos:{nombre}')).status_code, 403, nombre)
            self.entrar(self.vigilante)
            self.assertIn(self.client.get(reverse(f'accesos:{nombre}')).status_code, (302, 403), nombre)
            self.entrar(self.admin_accesos)
            self.assertEqual(self.client.get(reverse(f'accesos:{nombre}')).status_code, 200, nombre)

    def test_conductores_con_filtro_y_estado_de_licencia(self):
        Conductor.objects.create(tipo_documento=self.dni, numero_documento='07654321', apellido_paterno='ROJAS',
                                 nombres='ANA', licencia_numero='Q07654321', licencia_categoria='A-IIIb',
                                 licencia_vencimiento=HOY + timedelta(days=10), creado_por=self.despacho)
        datos = filas(libro(self.client.get(reverse('accesos:conductores_exportar')))['Datos'])
        self.assertEqual([d['Apellido paterno'] for d in datos], ['PÉREZ', 'ROJAS'])
        rojas = datos[1]
        self.assertEqual((rojas['Número documento'], rojas['Licencia categoría'], rojas['Estado de la licencia'],
                          rojas['Notificar']), ('07654321', 'A-IIIb', 'Por vencer', 'SI'))
        self.assertEqual(datos[0]['Estado de la licencia'], 'Sin licencia registrada')
        respuesta = self.client.get(reverse('accesos:conductores_exportar'), {'licencia': 'sin'})
        self.assertEqual([d['Apellido paterno'] for d in filas(libro(respuesta).active)], ['PÉREZ'])

    def test_conductores_se_corrige_y_se_vuelve_a_subir_en_actualizar_licencias(self):
        for numero, apellido in (('07654321', 'ROJAS'), ('07654322', 'SOTO')):
            Conductor.objects.create(tipo_documento=self.dni, numero_documento=numero, apellido_paterno=apellido,
                                     nombres='ANA', licencia_numero=f'Q{numero}', licencia_categoria='A-IIIb',
                                     licencia_vencimiento=date(2027, 1, 31), creado_por=self.despacho)
        descargado = libro(self.client.get(reverse('accesos:conductores_exportar'), {'licencia': ''}))
        hoja = descargado['Datos']
        encabezados = [c.value for c in hoja[1]]
        # Renueva la licencia de ROJAS en el archivo descargado.
        for fila in hoja.iter_rows(min_row=2):
            if fila[encabezados.index('Apellido paterno')].value == 'ROJAS':
                fila[encabezados.index('Licencia vencimiento')].value = date(2031, 1, 31)
        buffer = BytesIO()
        descargado.save(buffer)
        respuesta = self.previsualizar('licencias', SimpleUploadedFile('Conductores.xlsx', buffer.getvalue()))
        resultado = {f.valores['numero_documento']: f.estado for f in respuesta.context['filas']}
        self.assertEqual(resultado, {'12345678': carga_masiva.ERROR,  # PÉREZ no tiene licencia registrada
                                     '07654321': carga_masiva.OK, '07654322': carga_masiva.OMITIDA})
        self.confirmar('licencias')
        self.assertEqual(Conductor.objects.get(numero_documento='07654321').licencia_vencimiento, date(2031, 1, 31))


class DisponibilidadYAuditoriaExportarTests(PorteriaBase):

    def test_disponibilidad_con_averia_abierta_y_programaciones_por_reasignar(self):
        self.programar(15, 0, tipo='INTERNO')
        jornadas.registrar_no_disponibilidad(transporte_id=self.f7j.pk,
                                             tipo_id=TipoIncidencia.objects.get(nombre='Avería mecánica').pk,
                                             descripcion='Frenos', usuario=self.despacho, ahora=tiempo.ahora())
        self.programar(10, 0, fecha=MANANA, tipo='INTERNO')
        self.entrar(self.consulta)
        respuesta = self.client.get(reverse('accesos:flota_exportar'))
        datos = filas(libro(respuesta)['Incidencias'])
        self.assertEqual(len(datos), 1)
        self.assertEqual((datos[0]['Placa'], datos[0]['Estado'], datos[0]['Rutas canceladas'],
                          datos[0]['Programaciones por reasignar']), ('F7J-011', 'No disponible', 1, 1))
        self.assertContains(self.client.get(reverse('accesos:flota_disponibilidad')),
                            reverse('accesos:flota_exportar'))
        self.entrar(self.vigilante)
        self.assertEqual(self.client.get(reverse('accesos:flota_exportar')).status_code, 403)

    def test_auditoria_solo_administrador_con_filtros(self):
        servicios.cancelar(self.programar(10, 0), self.despacho, 'Cliente reprogramó', tiempo.ahora())
        self.entrar(self.despacho)
        self.assertEqual(self.client.get(reverse('accesos:auditoria_exportar')).status_code, 403)
        self.entrar(self.admin_accesos)
        respuesta = self.client.get(reverse('accesos:auditoria_exportar'), {'accion': 'CANCELAR'})
        datos = filas(libro(respuesta)['Auditoría'])
        self.assertEqual(len(datos), RegistroAuditoria.objects.filter(accion='CANCELAR').count())
        self.assertEqual((datos[0]['Motivo'], datos[0]['Usuario']), ('Cliente reprogramó', 'despacho'))
        self.assertContains(self.client.get(reverse('accesos:auditoria_lista'), {'accion': 'CANCELAR'}),
                            reverse('accesos:auditoria_exportar') + '?accion=CANCELAR')
