"""
Licencias de conducir del conductor: validación, edición por Despacho (solo licencia), alertas de la
campanita (30 días antes / 7 días después), desactivación automática, reactivación al renovar,
carga masiva y que nada de esto toque Planificación.
"""
from datetime import timedelta
from io import BytesIO, StringIO

from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from .. import licencias
from ..forms import ConductorForm
from ..models import ConfiguracionAccesos, Conductor, RegistroAuditoria
from .test_carga_masiva import CargaBase, excel
from .test_programacion import HOY, ProgramacionBase

COLUMNAS_LICENCIAS = ['tipo_documento', 'numero_documento', 'licencia_numero', 'licencia_categoria',
                      'licencia_vencimiento', 'licencia_especial_categoria', 'licencia_especial_vencimiento',
                      'notificar']


def dias(n):
    return HOY + timedelta(days=n)


class LicenciasBase(ProgramacionBase):

    def nuevo(self, numero, licencia=None, especial=None, **extra):
        """Conductor con licencia (días al vencimiento) y opcionalmente licencia especial A-IV."""
        datos = {'tipo_documento': self.dni, 'numero_documento': numero, 'apellido_paterno': f'AP{numero}',
                 'nombres': 'NOMBRE', 'creado_por': self.admin_accesos}
        if licencia is not None:
            datos.update(licencia_numero=f'Q{numero}', licencia_categoria='A-IIIb', licencia_vencimiento=dias(licencia))
        if especial is not None:
            datos.update(licencia_especial_categoria='A-IV', licencia_especial_vencimiento=dias(especial))
        datos.update(extra)
        return Conductor.objects.create(**datos)

    def form_datos(self, conductor, **cambios):
        datos = {'tipo_documento': conductor.tipo_documento_id, 'numero_documento': conductor.numero_documento,
                 'apellido_paterno': conductor.apellido_paterno, 'apellido_materno': conductor.apellido_materno,
                 'nombres': conductor.nombres, 'activo': 'on',
                 'licencia_numero': conductor.licencia_numero, 'licencia_categoria': conductor.licencia_categoria,
                 'licencia_vencimiento': conductor.licencia_vencimiento or '',
                 'licencia_especial_categoria': conductor.licencia_especial_categoria,
                 'licencia_especial_vencimiento': conductor.licencia_especial_vencimiento or ''}
        if conductor.notificar_vencimiento:
            datos['notificar_vencimiento'] = 'on'
        datos.update(cambios)
        return {k: v for k, v in datos.items() if v is not None}


class ValidacionTests(LicenciasBase):

    def test_licencia_completa_y_normalizada(self):
        form = ConductorForm(self.form_datos(self.conductor, licencia_numero='q-430 28991',
                                             licencia_categoria='A-IIIb', licencia_vencimiento='2029-06-07',
                                             licencia_especial_categoria='A-IV',
                                             licencia_especial_vencimiento='2027-12-12'),
                             instance=self.conductor)
        self.assertTrue(form.is_valid(), form.errors)
        conductor = form.save()
        self.assertEqual(conductor.licencia_numero, 'Q43028991')
        self.assertEqual(conductor.licencias(), [('Licencia', 'A-IIIb', dias(0).replace(2029, 6, 7)),
                                                 ('Licencia especial', 'A-IV', dias(0).replace(2027, 12, 12))])

    def test_sin_licencia_es_valido_y_las_notificaciones_vienen_activas(self):
        self.assertTrue(self.conductor.notificar_vencimiento)
        self.assertFalse(self.conductor.tiene_licencia)
        self.conductor.full_clean()

    def test_licencia_incompleta(self):
        casos = [
            ({'licencia_numero': 'Q12345678'}, {'licencia_categoria', 'licencia_vencimiento'}),
            ({'licencia_vencimiento': dias(100)}, {'licencia_numero', 'licencia_categoria'}),
            ({'licencia_especial_categoria': 'A-IV'}, {'licencia_especial_vencimiento', 'licencia_numero'}),
            ({'licencia_especial_vencimiento': dias(100)}, {'licencia_especial_categoria'}),
            ({'licencia_numero': 'Q1', 'licencia_categoria': 'A-IIb', 'licencia_vencimiento': dias(10)},
             {'licencia_numero'}),
        ]
        for campos, errores in casos:
            conductor = Conductor(tipo_documento=self.dni, numero_documento='55556666', apellido_paterno='X',
                                  nombres='Y', creado_por=self.admin_accesos, **campos)
            with self.subTest(campos=campos), self.assertRaises(ValidationError) as error:
                conductor.full_clean()
            self.assertTrue(errores <= set(error.exception.message_dict), error.exception.message_dict)

    def test_categoria_fuera_de_lista(self):
        form = ConductorForm(self.form_datos(self.conductor, licencia_numero='Q12345678', licencia_categoria='B-I',
                                             licencia_vencimiento='2029-01-01'), instance=self.conductor)
        self.assertFalse(form.is_valid())
        self.assertIn('licencia_categoria', form.errors)


class PantallasTests(LicenciasBase):

    def test_despacho_registra_conductor_con_licencia(self):
        self.entrar(self.despacho)
        self.assertContains(self.client.get(reverse('accesos:conductor_crear')), 'Notificar vencimientos')
        respuesta = self.client.post(reverse('accesos:conductor_crear'), {
            'tipo_documento': self.dni.pk, 'numero_documento': '43028991', 'apellido_paterno': 'Quispe',
            'nombres': 'Mario', 'activo': 'on', 'licencia_numero': 'Q43028991', 'licencia_categoria': 'A-IIIb',
            'licencia_vencimiento': '2029-06-07', 'licencia_especial_categoria': 'A-IV',
            'licencia_especial_vencimiento': '2027-12-12', 'notificar_vencimiento': 'on'})
        self.assertRedirects(respuesta, reverse('accesos:conductores_lista'))
        conductor = Conductor.objects.get(numero_documento='43028991')
        self.assertEqual((conductor.licencia_categoria, conductor.licencia_especial_categoria), ('A-IIIb', 'A-IV'))
        self.assertTrue(conductor.notificar_vencimiento)

    def test_registro_rapido_desde_programacion_mantiene_notificaciones(self):
        self.entrar(self.despacho)
        respuesta = self.client.post(reverse('accesos:conductor_crear_rapido'), {
            'co-tipo_documento': self.dni.pk, 'co-numero_documento': '77778888', 'co-apellido_paterno': 'Rojas',
            'co-nombres': 'Ana', 'co-activo': 'on'})
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        conductor = Conductor.objects.get(numero_documento='77778888')
        self.assertTrue(conductor.notificar_vencimiento)
        self.assertFalse(conductor.tiene_licencia)

    def test_despacho_edita_solo_la_licencia(self):
        self.entrar(self.despacho)
        url = reverse('accesos:conductor_editar', args=[self.conductor.pk])
        pagina = self.client.get(url)
        self.assertEqual(pagina.status_code, 200)
        self.assertTrue(pagina.context['solo_licencia'])
        respuesta = self.client.post(url, self.form_datos(
            self.conductor, apellido_paterno='CAMBIADO', numero_documento='99999999', activo=None,
            licencia_numero='Q12345678', licencia_categoria='A-IIIc', licencia_vencimiento='2028-01-15'))
        self.assertRedirects(respuesta, reverse('accesos:conductores_lista'))
        self.conductor.refresh_from_db()
        self.assertEqual(self.conductor.apellido_paterno, 'PÉREZ')       # identidad bloqueada para Despacho
        self.assertEqual(self.conductor.numero_documento, '12345678')
        self.assertTrue(self.conductor.activo)
        self.assertEqual(self.conductor.licencia_categoria, 'A-IIIc')
        registro = RegistroAuditoria.objects.filter(entidad='Conductor', objeto_id=self.conductor.pk).latest('id')
        self.assertEqual(registro.usuario, self.despacho)
        self.assertEqual(set(registro.datos_despues), {'licencia_numero', 'licencia_categoria',
                                                       'licencia_vencimiento', 'modificado_por_id'})

    def test_administrador_edita_todo(self):
        self.entrar(self.admin_accesos)
        url = reverse('accesos:conductor_editar', args=[self.conductor.pk])
        self.assertFalse(self.client.get(url).context['solo_licencia'])
        self.client.post(url, self.form_datos(self.conductor, apellido_paterno='Gómez'))
        self.conductor.refresh_from_db()
        self.assertEqual(self.conductor.apellido_paterno, 'GÓMEZ')

    def test_permisos_de_edicion(self):
        url = reverse('accesos:conductor_editar', args=[self.conductor.pk])
        for usuario in (self.consulta, self.vigilante, self.comercial, self.sin_grupo):
            self.entrar(usuario)
            self.assertEqual(self.client.get(url).status_code, 403, usuario)
        for usuario in (self.admin_accesos, self.despacho, self.planeamiento, self.superusuario):
            self.entrar(usuario)
            self.assertEqual(self.client.get(url).status_code, 200, usuario)

    def test_lista_con_estado_y_filtros(self):
        por_vencer = self.nuevo('20000001', licencia=10)
        vencida = self.nuevo('20000002', licencia=-3)
        vigente = self.nuevo('20000003', licencia=400)
        self.entrar(self.despacho)
        url = reverse('accesos:conductores_lista')
        pagina = self.client.get(url)
        self.assertContains(pagina, 'Por vencer')
        self.assertContains(pagina, 'Vencida')
        self.assertContains(pagina, 'Sin licencia registrada')
        filtrar = lambda valor: {c.pk for c in self.client.get(url, {'licencia': valor}).context['page_obj']}  # noqa: E731
        self.assertEqual(filtrar('alerta'), {por_vencer.pk, vencida.pk})
        self.assertEqual(filtrar('vencidas'), {vencida.pk})
        self.assertEqual(filtrar('sin'), {self.conductor.pk})
        self.assertEqual(len(filtrar('otra')), 4)  # valor inválido: sin filtro
        self.assertIn(vigente.pk, filtrar(''))
        self.assertEqual({c.pk for c in self.client.get(url, {'q': 'Q20000001'}).context['page_obj']},
                         {por_vencer.pk})


class AlertasTests(LicenciasBase):

    def test_ventana_de_30_dias_antes_y_7_despues(self):
        dentro = [self.nuevo(f'3000000{i}', licencia=d) for i, d in enumerate((30, 0, -7, 1))]
        fuera = [self.nuevo(f'3100000{i}', licencia=d) for i, d in enumerate((31, -8))]
        ids = {a.conductor.pk for a in licencias.alertas()}
        self.assertEqual(ids, {c.pk for c in dentro})
        self.assertFalse(ids & {c.pk for c in fuera})
        self.assertEqual([a.dias for a in licencias.alertas()], [-7, 0, 1, 30])  # la más urgente primero
        textos = {a.dias: a.texto for a in licencias.alertas()}
        self.assertEqual(textos, {-7: 'venció hace 7 días', 0: 'vence hoy', 1: 'vence mañana', 30: 'vence en 30 días'})

    def test_cada_licencia_alerta_por_separado(self):
        conductor = self.nuevo('40000001', licencia=200, especial=5)
        alertas = licencias.alertas()
        self.assertEqual([(a.conductor.pk, a.etiqueta, a.categoria) for a in alertas],
                         [(conductor.pk, 'Licencia especial', 'A-IV')])
        self.nuevo('40000002', licencia=3, especial=20)
        self.assertEqual(len(licencias.alertas()), 3)

    def test_sin_notificaciones_o_inactivo_no_alertan(self):
        self.nuevo('50000001', licencia=5, notificar_vencimiento=False)
        self.nuevo('50000002', licencia=5, activo=False)
        self.assertEqual(licencias.alertas(), [])

    def test_plazos_configurables(self):
        self.nuevo('50000003', licencia=40)
        config = ConfiguracionAccesos.obtener()
        config.aviso_licencia_dias = 45
        config.save()
        self.assertEqual(len(licencias.alertas()), 1)

    def test_estado(self):
        config = ConfiguracionAccesos.obtener()
        self.assertEqual(licencias.estado(dias(-1), HOY, config), licencias.VENCIDA)
        self.assertEqual(licencias.estado(dias(0), HOY, config), licencias.POR_VENCER)
        self.assertEqual(licencias.estado(dias(30), HOY, config), licencias.POR_VENCER)
        self.assertEqual(licencias.estado(dias(31), HOY, config), licencias.VIGENTE)
        self.assertIsNone(licencias.estado(None, HOY, config))


class DesactivacionAutomaticaTests(LicenciasBase):

    def test_se_desactiva_cuando_todas_las_licencias_pasaron_la_semana(self):
        vencida = self.nuevo('60000001', licencia=-8)
        ambas_vencidas = self.nuevo('60000002', licencia=-30, especial=-8)
        especial_vigente = self.nuevo('60000003', licencia=-30, especial=100)   # la A-IV aún no vence
        en_gracia = self.nuevo('60000004', licencia=-7)                          # justo el día 7: aún avisa
        sin_licencia = self.nuevo('60000005')
        self.assertEqual(licencias.revisar(), 2)
        estado = lambda c: Conductor.objects.get(pk=c.pk).notificar_vencimiento  # noqa: E731
        self.assertFalse(estado(vencida))
        self.assertFalse(estado(ambas_vencidas))
        self.assertTrue(estado(especial_vigente))
        self.assertTrue(estado(en_gracia))
        self.assertTrue(estado(sin_licencia))
        registro = RegistroAuditoria.objects.get(entidad='Conductor', objeto_id=vencida.pk, motivo__startswith='Desact')
        self.assertIsNone(registro.usuario)
        self.assertEqual(registro.autor, 'Sistema')
        self.assertEqual(registro.datos_antes, {'notificar_vencimiento': True})
        self.assertEqual(registro.datos_despues, {'notificar_vencimiento': False})
        self.assertIn('por Sistema', str(registro))
        self.assertEqual(licencias.revisar(), 0)  # idempotente

    def test_una_sola_revision_por_dia(self):
        config = ConfiguracionAccesos.obtener()
        self.nuevo('60000006', licencia=-10)
        self.assertEqual(licencias.revisar_si_corresponde(), 1)
        self.nuevo('60000007', licencia=-10)
        self.assertEqual(licencias.revisar_si_corresponde(), 0)  # ya se revisó hoy
        config.refresh_from_db()
        self.assertEqual(config.licencias_revisadas_el, HOY)
        self.mover_reloj(dias(1), 9)
        self.assertEqual(licencias.revisar_si_corresponde(), 1)  # al día siguiente vuelve a revisar

    def test_la_campanita_dispara_la_revision_diaria(self):
        conductor = self.nuevo('60000008', licencia=-10)
        self.entrar(self.despacho)
        self.client.get(reverse('accesos:inicio'))
        conductor.refresh_from_db()
        self.assertFalse(conductor.notificar_vencimiento)

    def test_comando(self):
        self.nuevo('60000009', licencia=-10)
        self.nuevo('60000010', licencia=12)
        salida = StringIO()
        call_command('revisar_licencias', stdout=salida)
        self.assertIn('Notificaciones desactivadas: 1. Alertas vigentes: 1.', salida.getvalue())
        self.assertEqual(ConfiguracionAccesos.obtener().licencias_revisadas_el, HOY)


class RenovacionTests(LicenciasBase):

    def setUp(self):
        super().setUp()
        self.vencido = self.nuevo('70000001', licencia=-20, notificar_vencimiento=False)
        self.url = reverse('accesos:conductor_editar', args=[self.vencido.pk])
        self.entrar(self.despacho)

    def test_al_renovar_se_reactivan_las_notificaciones(self):
        self.client.post(self.url, self.form_datos(self.vencido, licencia_vencimiento=dias(5 * 365).isoformat()))
        self.vencido.refresh_from_db()
        self.assertTrue(self.vencido.notificar_vencimiento)

    def test_si_las_apaga_al_renovar_se_respeta(self):
        self.vencido.notificar_vencimiento = True
        self.vencido.save()
        self.client.post(self.url, self.form_datos(self.vencido, licencia_vencimiento=dias(365).isoformat(),
                                                   notificar_vencimiento=None))
        self.vencido.refresh_from_db()
        self.assertFalse(self.vencido.notificar_vencimiento)

    def test_corregir_una_fecha_pasada_no_reactiva(self):
        self.client.post(self.url, self.form_datos(self.vencido, licencia_vencimiento=dias(-15).isoformat()))
        self.vencido.refresh_from_db()
        self.assertFalse(self.vencido.notificar_vencimiento)

    def test_renovar_la_especial_tambien_reactiva(self):
        self.client.post(self.url, self.form_datos(self.vencido, licencia_especial_categoria='A-IV',
                                                   licencia_especial_vencimiento=dias(90).isoformat()))
        self.vencido.refresh_from_db()
        self.assertTrue(self.vencido.notificar_vencimiento)


class CampanitaTests(LicenciasBase):

    def test_administrador_y_despacho_la_ven_con_el_conteo(self):
        self.nuevo('80000001', licencia=3)
        self.nuevo('80000002', licencia=-2)
        for usuario in (self.admin_accesos, self.despacho, self.planeamiento):
            self.entrar(usuario)
            pagina = self.client.get(reverse('accesos:inicio'))
            self.assertContains(pagina, 'id="campanitaBoton"')
            self.assertContains(pagina, '2 pendientes')
            self.assertContains(pagina, 'Venció hace 2 días')
            self.assertContains(pagina, 'Vence en 3 días')

    def test_sin_alertas(self):
        self.entrar(self.despacho)
        self.assertContains(self.client.get(reverse('accesos:inicio')), 'No hay licencias por vencer')

    def test_vigilancia_y_consulta_no_la_ven(self):
        self.nuevo('80000003', licencia=3)
        for usuario in (self.vigilante, self.consulta):
            self.entrar(usuario)
            self.assertNotContains(self.client.get(reverse('accesos:tablero_hoy')), 'campanitaBoton')

    def test_planificacion_no_consulta_licencias(self):
        """El context processor es perezoso: en Planificación no hay campanita ni revisión diaria."""
        self.nuevo('80000004', licencia=-10)
        self.entrar(self.despacho)
        pagina = self.client.get(reverse('dashboard'))
        self.assertEqual(pagina.status_code, 200)
        self.assertNotContains(pagina, 'campanitaBoton')
        self.assertIsNone(ConfiguracionAccesos.obtener().licencias_revisadas_el)

    def test_auditoria_muestra_sistema(self):
        self.nuevo('80000005', licencia=-10)
        licencias.revisar()
        self.entrar(self.admin_accesos)
        pagina = self.client.get(reverse('accesos:auditoria_lista'))
        self.assertContains(pagina, 'Sistema')
        self.assertContains(pagina, 'notificar_vencimiento: Sí → No')  # False no se muestra como vacío


class CargaLicenciasTests(CargaBase):

    def setUp(self):
        super().setUp()
        self.vencido = Conductor.objects.create(
            tipo_documento=self.dni, numero_documento='08123456', apellido_paterno='RAMOS', nombres='LUIS',
            licencia_numero='Q08123456', licencia_categoria='A-IIIa', licencia_vencimiento=dias(-20),
            licencia_especial_categoria='A-IV', licencia_especial_vencimiento=dias(-20),
            notificar_vencimiento=False, creado_por=self.admin_accesos)

    def test_conductores_con_licencia_desde_excel(self):
        archivo = excel(['tipo_documento', 'numero_documento', 'apellido_paterno', 'apellido_materno', 'nombres',
                         'licencia_numero', 'licencia_categoria', 'licencia_vencimiento',
                         'licencia_especial_categoria', 'licencia_especial_vencimiento'], [
            ['DNI', '43028991', 'Quispe', '', 'Mario', 'Q-43028991', 'A IIIb', '07/06/2029', 'a-iv', '12/12/2027'],
            ['DNI', '43028992', 'Sin', '', 'Licencia', '', '', '', '', ''],
            ['DNI', '43028993', 'Incompleta', '', 'X', 'Q43028993', 'A-IIIc', '', '', ''],
            ['DNI', '43028994', 'Mala', '', 'Categoria', 'Q43028994', 'B-IIa', '01/01/2030', '', ''],
        ])
        respuesta = self.previsualizar('conductores', archivo)
        self.assertEqual(self.estados(respuesta), [(2, 'OK'), (3, 'OK'), (4, 'ERROR'), (5, 'ERROR')])
        self.confirmar('conductores')
        mario = Conductor.objects.get(numero_documento='43028991')
        self.assertEqual((mario.licencia_numero, mario.licencia_categoria, mario.licencia_especial_categoria),
                         ('Q43028991', 'A-IIIb', 'A-IV'))
        self.assertEqual(mario.licencia_vencimiento.isoformat(), '2029-06-07')
        self.assertTrue(mario.notificar_vencimiento)
        self.assertFalse(Conductor.objects.get(numero_documento='43028992').tiene_licencia)

    def test_actualizar_licencias(self):
        archivo = excel(COLUMNAS_LICENCIAS, [
            ['DNI', 8123456, 'Q08123456', 'A-IIIb', '07/06/2031', '', '', ''],      # renueva; A-IV se conserva
            ['DNI', '12345678', 'Q12345678', 'A-IIb', '01/01/2030', 'A-IV', '01/03/2028', 'NO'],
            ['DNI', '99990000', 'Q99990000', 'A-IIb', '01/01/2030', '', '', ''],     # no existe
            ['DNI', '12345678', 'Q12345678', 'A-IIb', '01/01/2030', '', '', ''],     # repetido
            ['DNI', '12345678', 'Q12345678', 'A-IIb', '', '', '', 'quizas'],         # error
        ])
        respuesta = self.previsualizar('licencias', archivo)
        self.assertEqual(self.estados(respuesta), [(2, 'OK'), (3, 'OK'), (4, 'ERROR'), (5, 'ERROR'), (6, 'ERROR')])
        self.assertIn('no está registrado', ' '.join(respuesta.context['filas'][2].mensajes))
        self.assertEqual(Conductor.objects.get(pk=self.vencido.pk).licencia_categoria, 'A-IIIa')  # nada grabado aún
        resultado = self.confirmar('licencias')
        self.assertContains(resultado, 'Actualizada')
        self.vencido.refresh_from_db()
        self.assertEqual(self.vencido.licencia_categoria, 'A-IIIb')
        self.assertEqual(self.vencido.licencia_especial_categoria, 'A-IV')       # conservada
        self.assertEqual(self.vencido.licencia_especial_vencimiento, dias(-20))
        self.assertTrue(self.vencido.notificar_vencimiento)                      # renovó: se reactiva
        self.assertEqual(self.vencido.apellido_paterno, 'RAMOS')
        self.conductor.refresh_from_db()
        self.assertEqual(self.conductor.licencia_especial_categoria, 'A-IV')
        self.assertFalse(self.conductor.notificar_vencimiento)                   # el archivo dijo NO
        registro = RegistroAuditoria.objects.filter(entidad='Conductor', objeto_id=self.vencido.pk).latest('id')
        self.assertTrue(registro.motivo.startswith('Carga masiva:'))
        self.assertEqual(registro.usuario, self.despacho)

    def test_sin_cambios_se_omite(self):
        fila = ['DNI', '08123456', 'Q08123456', 'A-IIIa', dias(-20).strftime('%d/%m/%Y'), '', '', '']
        respuesta = self.previsualizar('licencias', excel(COLUMNAS_LICENCIAS, [fila]))
        self.assertEqual(self.estados(respuesta), [(2, 'OMITIDA')])

    def test_el_ejemplo_de_las_plantillas_se_puede_cargar(self):
        """La fila de la hoja «Ejemplo», copiada a «Datos», pasa la validación (fechas futuras, listas)."""
        Conductor.objects.filter(pk=self.vencido.pk).update(numero_documento='08123457')  # libera el DNI del ejemplo
        for tipo, esperado in (('conductores', 'OK'), ('licencias', 'ERROR')):  # licencias: el DNI no existe
            libro = load_workbook(BytesIO(self.client.get(reverse('accesos:carga_plantilla', args=[tipo])).content))
            for columna, celda in enumerate(libro['Ejemplo'][2], start=1):  # el usuario escribe desde la fila 2
                libro['Datos'].cell(row=2, column=columna, value=celda.value)
            buffer = BytesIO()
            libro.save(buffer)
            respuesta = self.previsualizar(tipo, SimpleUploadedFile(f'{tipo}.xlsx', buffer.getvalue()))
            self.assertEqual(self.estados(respuesta), [(2, esperado)], respuesta.context['filas'][0].mensajes)
        conductor = Conductor.objects.get(numero_documento='08123457')
        fila = ['DNI', '08123457', 'Q08123457', 'A-IIIb', libro['Ejemplo']['E2'].value, None, None, 'SI']
        for columna, valor in enumerate(fila, start=1):
            libro['Datos'].cell(row=2, column=columna, value=valor)
        buffer = BytesIO()
        libro.save(buffer)
        respuesta = self.previsualizar('licencias', SimpleUploadedFile('licencias.xlsx', buffer.getvalue()))
        self.assertEqual(self.estados(respuesta), [(2, 'OK')], respuesta.context['filas'][0].mensajes)
        self.confirmar('licencias')
        conductor.refresh_from_db()
        self.assertEqual(conductor.licencia_categoria, 'A-IIIb')
        self.assertTrue(conductor.notificar_vencimiento)

    def test_plantilla_de_licencias(self):
        respuesta = self.client.get(reverse('accesos:carga_plantilla', args=['licencias']))
        self.assertEqual(respuesta.status_code, 200)
        pagina = self.client.get(self.url('licencias'))
        self.assertContains(pagina, 'Actualizar licencias de conductores')
        self.assertContains(pagina, 'Volver a Conductores')
